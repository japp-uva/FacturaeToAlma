# SPDX-License-Identifier: GPL-2.0-only
#
# Conversor XSIG / Facturae a Excel de carga en Alma
# Copyright (C) 2026
#
# This program is free software; you can redistribute it and/or modify
# it under the terms of the GNU General Public License version 2 only,
# as published by the Free Software Foundation.

import os
import re
import copy
import html
import traceback
import logging
import unicodedata
import webbrowser
import subprocess
import platform
import xml.etree.ElementTree as ET

from decimal import Decimal, InvalidOperation
from dataclasses import dataclass, field
from datetime import datetime, date
from difflib import SequenceMatcher

import tkinter as tk
from tkinter import ttk, filedialog, messagebox

from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter


# ============================================================
# CONFIGURACIÓN GLOBAL
# ============================================================

APP_TITLE = "Universidad de Valladolid. Biblioteca → Conversor XSIG / Facturae → Excel Alma"
LOG_FILE = "facturae_alma.log"

PROJECT_REPOSITORY_URL = "https://github.com/japp-uva/FacturaeToAlma"
PROJECT_GUIDE_URL = "https://biblioguias.uva.es/facturae_to_alma"

UI_BG = "#f0f0f0"

TITLE_MATCH_MIN_SCORE = 0.78
TITLE_MATCH_MIN_MARGIN = 0.08

LINE_TYPE_OPTIONS = [
    "REGULAR",
    "OVERHEAD",
    "OTHER",
    "SHIPMENT",
    "DISCOUNT",
    "INSURANCE",
    "ADDITIONAL_CHARGES",
]

logging.basicConfig(
    filename=LOG_FILE,
    level=logging.INFO,
    format="%(asctime)s | %(levelname)s | %(message)s"
)

VENDOR_MAP = {
    "A1234567B": "JS_MAIN",
    "B98765432": "ELSEVIER_ES",
    "DEFAULT": "PROV_GENERICO"
}

DEFAULT_HINV_HEADERS = [
    "HINV", "Invoice Number", "Date", "Total amount", "Currency",
    "Vendor Account", "Payment method", "Note", "Use Pro Rata",
    "Shipment amount", "Overhead amount", "Insurance amount",
    "Discount amount", "VAT Percentage", "VAT Amount", "Inclusive",
    "Expended From Fund", "PrePaid", "Payment status",
    "Voucher number", "Voucher date", "Voucher Amount",
    "Invoice Reference Number", "VAT In Invoice Line Level", "Report TAX"
]

DEFAULT_HIL_HEADERS = [
    "HIL", "Line type", "Line Number", "PO Line", "Vendor Ref Number",
    "ISSN", "ISBN", "Title", "Price", "Quantity",
    "Reporting Code", "Secondary Reporting Code", "Tertiary Reporting Code",
    "Fourth Reporting Code", "Fifth Reporting Code",
    "Start subs date", "End subs date", "Note",
    "Fund and percent", "Fund and percent",
    "VAT Amount", "VAT Percentage"
]


# ============================================================
# STOPWORDS INTEGRADAS
# ============================================================

STOPWORDS = {
    # Español
    "el", "la", "los", "las", "un", "una", "unos", "unas",
    "a", "ante", "bajo", "con", "contra", "de", "desde", "en",
    "entre", "hacia", "hasta", "para", "por", "segun", "según",
    "sin", "sobre", "tras",
    "y", "e", "ni", "que", "pero", "aunque", "sino", "porque", "si", "sí",
    "me", "te", "se", "nos", "mi", "mí", "tu", "tú", "su",
    "este", "esta", "esto",

    # Inglés
    "a", "an", "the",
    "in", "on", "at", "by", "for", "with", "about", "from", "to",
    "of", "under", "over",
    "and", "but", "or", "so", "because", "although", "yet",
    "i", "you", "he", "she", "it", "we", "they", "this", "that", "there",

    # Francés
    "le", "la", "les", "l", "un", "une", "des",
    "a", "à", "de", "dans", "en", "sur", "avec", "pour", "par", "sans", "chez",
    "et", "ou", "ni", "mais", "donc", "car", "or",
    "je", "tu", "il", "elle", "nous", "vous", "ce", "se", "y",

    # Alemán
    "der", "die", "das", "dem", "den", "des", "ein", "eine", "einer",
    "in", "auf", "an", "zu", "mit", "von", "fur", "für", "bei",
    "durch", "aus", "vor",
    "und", "oder", "aber", "denn", "weil", "obwohl", "sondern",
    "ich", "du", "er", "sie", "es", "wir", "ihr", "sich",

    # Italiano
    "il", "lo", "la", "i", "gli", "le", "un", "uno", "una",
    "di", "a", "da", "in", "con", "su", "per", "tra", "fra",
    "e", "ed", "ma", "o", "oppure", "perche", "perché", "anche",
    "mi", "ti", "si", "ci", "io", "tu", "lui", "lei", "noi", "voi",

    # Catalán
    "el", "la", "els", "les", "l", "un", "una", "uns", "unes",
    "a", "de", "en", "amb", "per", "sense", "sota", "contra",
    "i", "o", "pero", "però", "perque", "perquè", "si", "ni", "doncs",
    "jo", "tu", "ell", "ella", "nosaltres", "vosaltres", "es", "em", "et",

    # Gallego
    "o", "a", "os", "as", "un", "unha", "uns", "unhas",
    "ante", "baixo", "con", "contra", "de", "desde", "en", "entre",
    "para", "por", "segundo", "sen", "sobre", "tras",
    "e", "ou", "pero", "senon", "senón", "mais", "porque", "se", "nin", "que",
    "eu", "ti", "el", "ela", "nos", "nós", "vos", "vós", "me", "te",

    # Euskera
    "a", "ak", "ari", "ariak", "bat", "batzuk",
    "baino", "gabe", "gora", "behera", "arte", "bila", "kontra", "zehar", "aurka",
    "eta", "edo", "baina", "denez", "beraz", "orduan", "baita", "ezta", "zein",
    "ni", "hi", "hura", "gu", "zu", "zuek", "haiek", "hau", "hori",

    # Portugués
    "o", "a", "os", "as", "um", "uma", "uns", "umas",
    "com", "contra", "de", "desde", "em", "entre", "para", "por",
    "sem", "sob", "sobre", "tras", "trás", "ate", "até",
    "e", "mas", "ou", "porque", "se", "nem", "que", "embora",
    "contudo", "pois",
    "eu", "tu", "ele", "ela", "nos", "nós", "vos", "vós",
    "me", "te", "este", "esta", "isto", "esse", "essa", "isso",
}

STOPWORD_PHRASES = {
    "per a",
    "des de",
    "cap a",
    "fins a",
    "encara que",
    "ainda que",
    "aínda que",
}


# ============================================================
# MODELOS
# ============================================================

@dataclass
class AlmaPOLRecord:
    po_line: str
    title: str = None
    reporting_code: str = None
    secondary_reporting_code: str = None
    tertiary_reporting_code: str = None
    fourth_reporting_code: str = None
    fifth_reporting_code: str = None
    start_subs_date: datetime = None
    isbn_raw: str = None
    issn_raw: str = None
    funds: list = field(default_factory=list)
    normalized_title_words: list = field(default_factory=list)
    normalized_title_text: str = ""


@dataclass
class ValidationIssue:
    issue_type: str
    invoice_line_number: int
    isbn: str
    issn: str
    title: str
    candidates: list
    message: str


@dataclass
class InvoiceLineData:
    line_number: int
    po_line: str = None
    vendor_ref: str = None
    issn: str = None
    isbn: str = None
    title: str = None

    price: Decimal = Decimal("0.00")
    quantity: Decimal = Decimal("1.00")

    start_date: datetime = None
    end_date: datetime = None
    note: str = None

    vat_amount: Decimal = None
    vat_percentage: Decimal = None

    reporting_code: str = None
    secondary_reporting_code: str = None
    tertiary_reporting_code: str = None
    fourth_reporting_code: str = None
    fifth_reporting_code: str = None

    funds: list = field(default_factory=list)

    match_status: str = None
    match_message: str = None
    match_candidates: list = field(default_factory=list)


@dataclass
class InvoiceData:
    invoice_number: str
    date: datetime
    total_amount: Decimal
    currency: str
    vendor_account: str
    note: str

    vat_amount: Decimal = None
    inclusive: bool = False

    lines: list = field(default_factory=list)


# ============================================================
# UTILIDADES
# ============================================================

def clean_cell(value):
    if value is None:
        return None

    text = str(value).replace("\xa0", " ").strip()

    if not text or text.lower() == "nan":
        return None

    return text


def alma_empty_to_none(value):
    value = clean_cell(value)

    if value in (None, "", "-1"):
        return None

    return value


def safe_decimal(value, default=Decimal("0.00")):
    if value is None:
        return default

    raw = str(value).strip()

    if not raw:
        return default

    raw = raw.replace(" ", "")

    if "." in raw and "," in raw:
        if raw.rfind(",") > raw.rfind("."):
            raw = raw.replace(".", "").replace(",", ".")
        else:
            raw = raw.replace(",", "")
    elif "," in raw and "." not in raw:
        raw = raw.replace(",", ".")

    try:
        return Decimal(raw)
    except InvalidOperation:
        logging.warning(f"No se pudo convertir a Decimal: {value!r}")
        return default


def decimal_or_none(value):
    if value is None:
        return None

    text = str(value).strip()

    if not text:
        return None

    return safe_decimal(text)


def decimal_to_float_or_none(value):
    if value is None:
        return None

    return float(value)


def parse_date_for_excel(value):
    if value is None:
        return None

    if isinstance(value, datetime):
        return value

    if isinstance(value, date):
        return datetime(value.year, value.month, value.day)

    text = clean_cell(value)

    if not text:
        return None

    text = text[:10]

    possible_formats = [
        "%Y-%m-%d",
        "%d-%m-%Y",
        "%d/%m/%Y",
        "%m/%d/%Y",
        "%Y/%m/%d",
    ]

    for fmt in possible_formats:
        try:
            return datetime.strptime(text, fmt)
        except ValueError:
            pass

    logging.warning(f"No se pudo convertir la fecha: {value!r}")
    return None


def normalize_identifier(value):
    if not value:
        return ""

    value = str(value).replace("\xa0", " ").upper()
    return re.sub(r"[^0-9X]", "", value)


def split_identifiers(value):
    value = clean_cell(value)

    if not value:
        return []

    result = []

    for part in value.split(";"):
        normalized = normalize_identifier(part)

        if normalized:
            result.append(normalized)

    return result


def is_valid_isbn10(code):
    code = normalize_identifier(code)

    if not re.fullmatch(r"\d{9}[\dX]", code):
        return False

    total = 0

    for i, char in enumerate(code):
        value = 10 if char == "X" else int(char)
        total += (10 - i) * value

    return total % 11 == 0


def is_valid_isbn13(code):
    code = normalize_identifier(code)

    if not re.fullmatch(r"\d{13}", code):
        return False

    total = 0

    for i, char in enumerate(code[:12]):
        value = int(char)
        total += value if i % 2 == 0 else value * 3

    check = (10 - (total % 10)) % 10
    return check == int(code[-1])


def is_valid_issn(code):
    code = normalize_identifier(code)

    if not re.fullmatch(r"\d{7}[\dX]", code):
        return False

    total = 0

    for i in range(7):
        total += int(code[i]) * (8 - i)

    check = (11 - (total % 11)) % 11
    expected = "X" if check == 10 else str(check)

    return code[-1] == expected


def looks_like_issn(code):
    clean = normalize_identifier(code)
    return bool(re.fullmatch(r"\d{7}[\dX]", clean))


def looks_like_isbn(code):
    clean = normalize_identifier(code)
    return bool(re.fullmatch(r"\d{10}|\d{13}", clean))


def extract_identifiers_from_text(text):
    """
    Busca ISBN/ISSN dentro de un texto.
    Devuelve:
        isbn, issn, cleaned_text
    """
    if not text:
        return None, None, text

    original_text = str(text)

    isbn = None
    issn = None
    matched_fragments = []

    candidate_pattern = re.compile(r"[0-9Xx][0-9Xx\-\s]{6,25}[0-9Xx]")

    for match in candidate_pattern.finditer(original_text):
        fragment = match.group(0)
        normalized = normalize_identifier(fragment)

        if not isbn:
            if len(normalized) == 13 and is_valid_isbn13(normalized):
                isbn = normalized
                matched_fragments.append(fragment)
                continue

            if len(normalized) == 10 and is_valid_isbn10(normalized):
                isbn = normalized
                matched_fragments.append(fragment)
                continue

        if not issn:
            if len(normalized) == 8 and is_valid_issn(normalized):
                issn = normalized
                matched_fragments.append(fragment)
                continue

    cleaned_text = original_text

    for fragment in matched_fragments:
        cleaned_text = cleaned_text.replace(fragment, " ")

    cleaned_text = re.sub(r"\s*/\s*$", "", cleaned_text)
    cleaned_text = re.sub(r"\s*[-–—:;,/]\s*$", "", cleaned_text)
    cleaned_text = re.sub(r"\s+", " ", cleaned_text).strip()

    return isbn, issn, cleaned_text


def classify_article_code(code):
    if not code:
        return None, None, None

    isbn, issn, _ = extract_identifiers_from_text(code)

    if isbn:
        return None, isbn, None

    if issn:
        return issn, None, None

    if looks_like_issn(code):
        return code, None, None

    if looks_like_isbn(code):
        return None, code, None

    return None, None, code


def normalize_header(header):
    if header is None:
        return ""

    return " ".join(str(header).strip().split())


def get_vendor_account(nif, corporate_name=None):
    nif = nif.strip() if nif else ""

    if nif and nif in VENDOR_MAP:
        return VENDOR_MAP[nif]

    if corporate_name:
        return corporate_name.strip()

    return VENDOR_MAP["DEFAULT"]


def build_validation_output_path(output_path):
    base, _ = os.path.splitext(output_path)
    return f"{base}_validacion.xlsx"


def remove_accents(text):
    if not text:
        return ""

    normalized = unicodedata.normalize("NFD", text)
    return "".join(
        char for char in normalized
        if unicodedata.category(char) != "Mn"
    )


def normalize_title_for_matching(title):
    if not title:
        return []

    text = html.unescape(str(title))
    text = remove_accents(text.lower())

    text = text.replace("’", "'")
    text = text.replace("´", "'")

    for phrase in STOPWORD_PHRASES:
        normalized_phrase = remove_accents(phrase.lower())
        normalized_phrase = re.sub(r"[^a-z0-9]+", " ", normalized_phrase).strip()

        if normalized_phrase:
            text = re.sub(rf"\b{re.escape(normalized_phrase)}\b", " ", text)

    text = re.sub(r"[^a-z0-9]+", " ", text)

    words = text.split()
    result = []

    for word in words:
        if len(word) <= 1:
            continue

        if word in STOPWORDS:
            continue

        result.append(word)

    return result


def title_similarity_score(invoice_title, record):
    invoice_words = normalize_title_for_matching(invoice_title)
    alma_words = record.normalized_title_words

    if not invoice_words or not alma_words:
        return 0.0

    invoice_set = set(invoice_words)
    alma_set = set(alma_words)

    common = invoice_set.intersection(alma_set)
    word_score = len(common) / len(invoice_set)

    invoice_clean = " ".join(invoice_words)
    alma_clean = record.normalized_title_text

    sequence_score = SequenceMatcher(
        None,
        invoice_clean,
        alma_clean
    ).ratio()

    return (word_score * 0.75) + (sequence_score * 0.25)


def build_conversion_summary(invoices, validation_issues):
    summary = {
        "total_lines": 0,
        "match_isbn": 0,
        "match_issn": 0,
        "match_title": 0,
        "duplicate_isbn": 0,
        "duplicate_issn": 0,
        "duplicate_title": 0,
        "ambiguous_title": 0,
        "no_match": 0,
        "validation_issues": len(validation_issues),
        "lines_without_po_line": 0,
    }

    for invoice in invoices:
        for line in invoice.lines:
            summary["total_lines"] += 1

            status = line.match_status or ""

            if status == "MATCH_ISBN":
                summary["match_isbn"] += 1
            elif status == "MATCH_ISSN":
                summary["match_issn"] += 1
            elif status == "MATCH_TITLE":
                summary["match_title"] += 1
            elif status == "DUPLICATE_ISBN":
                summary["duplicate_isbn"] += 1
            elif status == "DUPLICATE_ISSN":
                summary["duplicate_issn"] += 1
            elif status == "DUPLICATE_TITLE":
                summary["duplicate_title"] += 1
            elif status == "AMBIGUOUS_TITLE":
                summary["ambiguous_title"] += 1
            elif status == "NO_MATCH":
                summary["no_match"] += 1

            if not line.po_line:
                summary["lines_without_po_line"] += 1

    return summary


# ============================================================
# LECTOR DEL EXCEL DE ALMA
# ============================================================

class AlmaPOLineLookup:
    REQUIRED_COLUMNS = [
        "PO Line",
        "Reporting Code",
        "Secondary Reporting Code",
        "Tertiary Reporting Code",
        "Start subs date",
        "Title",
        "ISSN",
        "ISBN"
    ]

    def __init__(self, alma_excel_path):
        self.alma_excel_path = alma_excel_path
        self.records = []
        self.isbn_index = {}
        self.issn_index = {}

    def load(self):
        try:
            wb = load_workbook(self.alma_excel_path, data_only=True)
            ws = wb.active
        except Exception as e:
            logging.exception("Error cargando Excel de Alma")
            raise ValueError(f"No se pudo leer el Excel de Alma con PO Lines: {e}")

        headers = self._read_headers(ws)
        self._validate_headers(headers)

        col_map = {}
        col_map_multi = {}

        for idx, header in enumerate(headers, start=1):
            normalized = normalize_header(header)

            if normalized not in col_map:
                col_map[normalized] = idx

            col_map_multi.setdefault(normalized, []).append(idx)

        fund_columns = col_map_multi.get("Fund and percent", [])

        for row_idx in range(2, ws.max_row + 1):
            po_line = clean_cell(ws.cell(row=row_idx, column=col_map["PO Line"]).value)

            if not po_line:
                continue

            title = clean_cell(ws.cell(row=row_idx, column=col_map["Title"]).value)

            title_words = normalize_title_for_matching(title)
            title_text = " ".join(title_words)

            funds = []

            for fund_col in fund_columns:
                fund_value = alma_empty_to_none(ws.cell(row=row_idx, column=fund_col).value)

                if fund_value:
                    funds.append(fund_value)

            record = AlmaPOLRecord(
                po_line=po_line,
                reporting_code=alma_empty_to_none(
                    ws.cell(row=row_idx, column=col_map["Reporting Code"]).value
                ),
                secondary_reporting_code=alma_empty_to_none(
                    ws.cell(row=row_idx, column=col_map["Secondary Reporting Code"]).value
                ),
                tertiary_reporting_code=alma_empty_to_none(
                    ws.cell(row=row_idx, column=col_map["Tertiary Reporting Code"]).value
                ),
                fourth_reporting_code=alma_empty_to_none(
                    ws.cell(row=row_idx, column=col_map["Fourth Reporting Code"]).value
                ) if "Fourth Reporting Code" in col_map else None,
                fifth_reporting_code=alma_empty_to_none(
                    ws.cell(row=row_idx, column=col_map["Fifth Reporting Code"]).value
                ) if "Fifth Reporting Code" in col_map else None,
                start_subs_date=parse_date_for_excel(
                    ws.cell(row=row_idx, column=col_map["Start subs date"]).value
                ),
                title=title,
                issn_raw=clean_cell(ws.cell(row=row_idx, column=col_map["ISSN"]).value),
                isbn_raw=clean_cell(ws.cell(row=row_idx, column=col_map["ISBN"]).value),
                funds=funds,
                normalized_title_words=title_words,
                normalized_title_text=title_text
            )

            self.records.append(record)

            for isbn in split_identifiers(record.isbn_raw):
                self.isbn_index.setdefault(isbn, []).append(record)

            for issn in split_identifiers(record.issn_raw):
                self.issn_index.setdefault(issn, []).append(record)

        logging.info(
            f"Excel Alma cargado: {len(self.records)} registros, "
            f"{len(self.isbn_index)} ISBN indexados, "
            f"{len(self.issn_index)} ISSN indexados"
        )

    def _read_headers(self, ws):
        return [
            clean_cell(cell.value) or ""
            for cell in ws[1]
        ]

    def _validate_headers(self, headers):
        normalized = {normalize_header(header) for header in headers}
        missing = [col for col in self.REQUIRED_COLUMNS if col not in normalized]

        if missing:
            raise ValueError(
                "El Excel de Alma no contiene las columnas esperadas. "
                f"Faltan: {', '.join(missing)}"
            )

    def find_by_isbn(self, isbn):
        normalized = normalize_identifier(isbn)

        if not normalized:
            return []

        return self.isbn_index.get(normalized, [])

    def find_by_issn(self, issn):
        normalized = normalize_identifier(issn)

        if not normalized:
            return []

        return self.issn_index.get(normalized, [])

    def find_by_title(self, title):
        if not title:
            return {
                "status": "NO_MATCH",
                "candidates": [],
                "scores": []
            }

        scored = []

        for record in self.records:
            if not record.normalized_title_words:
                continue

            score = title_similarity_score(title, record)

            if score >= TITLE_MATCH_MIN_SCORE:
                scored.append((record, score))

        scored.sort(key=lambda item: item[1], reverse=True)

        if not scored:
            return {
                "status": "NO_MATCH",
                "candidates": [],
                "scores": []
            }

        best_record, best_score = scored[0]

        if len(scored) == 1:
            return {
                "status": "MATCH",
                "candidates": [best_record],
                "scores": [best_score]
            }

        second_score = scored[1][1]

        if best_score - second_score >= TITLE_MATCH_MIN_MARGIN:
            return {
                "status": "MATCH",
                "candidates": [best_record],
                "scores": [best_score]
            }

        ambiguous_records = [
            record for record, score in scored
            if best_score - score < TITLE_MATCH_MIN_MARGIN
        ]

        ambiguous_scores = [
            score for record, score in scored
            if best_score - score < TITLE_MATCH_MIN_MARGIN
        ]

        return {
            "status": "AMBIGUOUS",
            "candidates": ambiguous_records,
            "scores": ambiguous_scores
        }


# ============================================================
# PARSER XSIG / XML / TXT
# ============================================================

class FacturaeParser:
    def __init__(self, invoice_path):
        self.invoice_path = invoice_path
        self.tree = None
        self.root = None

    def load(self):
        try:
            self.tree = ET.parse(self.invoice_path)
            self.root = self.tree.getroot()
            logging.info(f"Archivo de factura cargado correctamente: {self.invoice_path}")
            return

        except Exception as first_error:
            logging.warning(
                f"No se pudo parsear directamente. "
                f"Se intentará lectura alternativa. Error inicial: {first_error}"
            )

        encodings = ["utf-8-sig", "utf-8", "iso-8859-1", "latin-1"]
        last_error = None

        for enc in encodings:
            try:
                with open(self.invoice_path, "r", encoding=enc) as f:
                    content = f.read()

                content = content.strip()

                xml_start = content.find("<")

                if xml_start > 0:
                    content = content[xml_start:]

                self.root = ET.fromstring(content)
                self.tree = ET.ElementTree(self.root)

                logging.info(
                    f"Archivo de factura cargado mediante lectura alternativa "
                    f"con codificación {enc}: {self.invoice_path}"
                )
                return

            except Exception as e:
                last_error = e

        logging.exception("Error cargando archivo de factura")
        raise ValueError(
            "No se pudo leer o parsear el archivo de factura. "
            "Debe ser un XML válido, aunque tenga extensión .xsig, .xml o .txt. "
            f"Último error: {last_error}"
        )

    def find_text(self, element, tag, default=""):
        if element is None:
            return default

        found = element.find(f".//{{*}}{tag}")

        if found is not None and found.text:
            return found.text.strip()

        return default

    def find_path_text(self, element, path, default=""):
        if element is None:
            return default

        parts = path.split("/")
        xpath = "./" + "/".join(f"{{*}}{part}" for part in parts)

        found = element.find(xpath)

        if found is not None and found.text:
            return found.text.strip()

        return default

    def parse(self):
        if self.root is None:
            self.load()

        invoices_xml = self.root.findall(".//{*}Invoice")
        logging.info(f"Número de nodos Invoice detectados: {len(invoices_xml)}")

        if not invoices_xml:
            raise ValueError(
                "No se encontraron elementos <Invoice> en el archivo. "
                "Revisa si el XSIG/XML/TXT contiene una factura Facturae válida."
            )

        return [self._parse_invoice(inv) for inv in invoices_xml]

    def _parse_invoice(self, inv):
        seller_party = self.root.find(".//{*}SellerParty")

        nif = self.find_text(seller_party, "TaxIdentificationNumber", "")
        corporate_name = self.find_text(seller_party, "CorporateName", "")

        vendor_account = get_vendor_account(nif, corporate_name)

        invoice_number = self.find_text(inv, "InvoiceNumber")
        issue_date = parse_date_for_excel(self.find_text(inv, "IssueDate"))
        total_amount = safe_decimal(self.find_text(inv, "InvoiceTotal"))
        currency = self.find_text(inv, "InvoiceCurrencyCode", "EUR")
        note = self.find_text(inv, "InvoiceDescription") or None

        vat_amount = decimal_or_none(
            self.find_path_text(
                inv,
                "TaxesOutputs/Tax/TaxAmount/TotalAmount"
            )
        )

        invoice_data = InvoiceData(
            invoice_number=invoice_number,
            date=issue_date,
            total_amount=total_amount,
            currency=currency,
            vendor_account=vendor_account,
            note=note,
            vat_amount=vat_amount,
            inclusive=False,
            lines=[]
        )

        lines_xml = inv.findall(".//{*}InvoiceLine")
        logging.info(f"Factura {invoice_number}: {len(lines_xml)} líneas")

        for idx, line_xml in enumerate(lines_xml, start=1):
            invoice_data.lines.append(self._parse_line(line_xml, idx))

        return invoice_data

    def _parse_line(self, line, idx):
        seq_str = self.find_text(line, "SequenceNumber", str(idx))

        try:
            seq_num = int(seq_str)
        except ValueError:
            seq_num = idx

        raw_title = self.find_text(line, "ItemDescription") or None
        desc_isbn, desc_issn, cleaned_title = extract_identifiers_from_text(raw_title)

        title = cleaned_title or raw_title

        quantity = safe_decimal(
            self.find_text(line, "Quantity", "1"),
            Decimal("1.00")
        )

        gross_amount_text = self.find_path_text(line, "GrossAmount")
        unit_price_text = self.find_text(line, "UnitPriceWithoutTax")

        if gross_amount_text:
            price = safe_decimal(gross_amount_text, Decimal("0.00"))
        else:
            price = safe_decimal(unit_price_text, Decimal("0.00"))

        article_code = self.find_text(line, "ArticleCode") or None

        issn, isbn, vendor_ref = classify_article_code(article_code)

        if not isbn and desc_isbn:
            isbn = desc_isbn

        if not issn and desc_issn:
            issn = desc_issn

        po_line = None

        vat_amount_line = decimal_or_none(
            self.find_path_text(
                line,
                "TaxesOutputs/Tax/TaxAmount/TotalAmount"
            )
        )

        vat_percentage_line = decimal_or_none(
            self.find_path_text(
                line,
                "TaxesOutputs/Tax/TaxRate"
            )
        )

        return InvoiceLineData(
            line_number=seq_num,
            po_line=po_line,
            vendor_ref=vendor_ref,
            issn=issn,
            isbn=isbn,
            title=title,
            price=price,
            quantity=quantity,
            vat_amount=vat_amount_line,
            vat_percentage=vat_percentage_line
        )


# ============================================================
# ENRIQUECIMIENTO CON DATOS DE ALMA
# ============================================================

class AlmaLineEnricher:
    def __init__(self, lookup):
        self.lookup = lookup
        self.validation_issues = []

    def enrich(self, invoices):
        for invoice in invoices:
            for line in invoice.lines:
                self._enrich_line(line)

        return self.validation_issues

    def _clear_alma_fields(self, line):
        line.po_line = None
        line.reporting_code = None
        line.secondary_reporting_code = None
        line.tertiary_reporting_code = None
        line.fourth_reporting_code = None
        line.fifth_reporting_code = None
        line.funds = []

    def _enrich_line(self, line):
        if line.isbn:
            candidates = self.lookup.find_by_isbn(line.isbn)

            if candidates:
                self._resolve_candidates(line, candidates, "ISBN")
                return

        if line.issn:
            candidates = self.lookup.find_by_issn(line.issn)

            if candidates:
                self._resolve_candidates(line, candidates, "ISSN")
                return

        if line.title:
            result = self.lookup.find_by_title(line.title)

            if result["status"] == "MATCH":
                self._resolve_candidates(line, result["candidates"], "TITLE")
                return

            if result["status"] == "AMBIGUOUS":
                candidate_po_lines = [
                    record.po_line for record in result["candidates"]
                ]

                score_text = "; ".join(
                    f"{record.po_line}: {score:.3f}"
                    for record, score in zip(result["candidates"], result["scores"])
                )

                self._clear_alma_fields(line)

                line.match_status = "AMBIGUOUS_TITLE"
                line.match_message = "Varias coincidencias posibles por título"
                line.match_candidates = candidate_po_lines

                self.validation_issues.append(
                    ValidationIssue(
                        issue_type="TÍTULO AMBIGUO",
                        invoice_line_number=line.line_number,
                        isbn=line.isbn,
                        issn=line.issn,
                        title=line.title,
                        candidates=candidate_po_lines,
                        message=(
                            "Varias PO Lines candidatas por título. "
                            f"Puntuaciones: {score_text}"
                        )
                    )
                )
                return

        self._clear_alma_fields(line)

        line.match_status = "NO_MATCH"
        line.match_message = "Sin coincidencia en Alma por ISBN, ISSN ni título"

        self.validation_issues.append(
            ValidationIssue(
                issue_type="SIN COINCIDENCIA",
                invoice_line_number=line.line_number,
                isbn=line.isbn,
                issn=line.issn,
                title=line.title,
                candidates=[],
                message=(
                    "No se encontró ninguna PO Line por ISBN, ISSN ni título. "
                    "Puede tratarse de una línea ya cerrada/cargada en Alma, "
                    "de una línea que no está incluida en el listado de líneas abiertas "
                    "o de un título que requiere revisión manual."
                )
            )
        )

    def _resolve_candidates(self, line, candidates, match_type):
        unique_by_po = {}

        for candidate in candidates:
            unique_by_po[candidate.po_line] = candidate

        unique_candidates = list(unique_by_po.values())

        if len(unique_candidates) == 1:
            record = unique_candidates[0]

            line.po_line = record.po_line
            line.reporting_code = record.reporting_code
            line.secondary_reporting_code = record.secondary_reporting_code
            line.tertiary_reporting_code = record.tertiary_reporting_code
            line.fourth_reporting_code = record.fourth_reporting_code
            line.fifth_reporting_code = record.fifth_reporting_code
            line.funds = list(record.funds or [])

            if record.title:
                line.title = record.title

            if record.start_subs_date:
                line.start_date = record.start_subs_date

            line.match_status = f"MATCH_{match_type}"
            line.match_message = f"Coincidencia única por {match_type}"
            line.match_candidates = [record.po_line]

            return

        candidate_po_lines = sorted(record.po_line for record in unique_candidates)

        self._clear_alma_fields(line)

        line.match_status = f"DUPLICATE_{match_type}"
        line.match_message = f"Duplicado por {match_type}. Revisión manual necesaria."
        line.match_candidates = candidate_po_lines

        self.validation_issues.append(
            ValidationIssue(
                issue_type=f"DUPLICADO {match_type}",
                invoice_line_number=line.line_number,
                isbn=line.isbn,
                issn=line.issn,
                title=line.title,
                candidates=candidate_po_lines,
                message=(
                    "El identificador tiene varias PO Lines candidatas. "
                    "No se ha rellenado PO Line automáticamente."
                )
            )
        )


# ============================================================
# EXPORTADOR EXCEL ALMA
# ============================================================

class AlmaExcelExporter:
    DATE_FORMAT = "mm/dd/yyyy"

    def __init__(self, output_path, template_path=None):
        self.output_path = output_path
        self.template_path = template_path

        self.workbook = None
        self.worksheet = None

        self.hinv_headers = []
        self.hil_headers = []

        self.hinv_style = None
        self.inv_style = None
        self.hil_style = None
        self.il_style = None

        self.column_widths = {}

    def create_or_load_workbook(self):
        if self.template_path and os.path.exists(self.template_path):
            self._load_from_template()
        else:
            self._create_without_template()

    def _load_from_template(self):
        template_wb = load_workbook(self.template_path)
        template_ws = template_wb.active

        hinv_row = self._find_marker_row(template_ws, "HINV")
        inv_row = hinv_row + 1 if hinv_row else 2

        hil_row = self._find_marker_row(template_ws, "HIL")
        il_row = hil_row + 1 if hil_row else 4

        if not hinv_row or not hil_row:
            raise ValueError(
                "La plantilla seleccionada no contiene filas HINV y HIL reconocibles."
            )

        self.hinv_headers = self._read_row_values(template_ws, hinv_row)
        self.hil_headers = self._read_row_values(template_ws, hil_row)

        max_cols = max(len(self.hinv_headers), len(self.hil_headers), 25)

        self.hinv_style = self._capture_row_style(template_ws, hinv_row, max_cols)
        self.inv_style = self._capture_row_style(template_ws, inv_row, max_cols)
        self.hil_style = self._capture_row_style(template_ws, hil_row, max_cols)
        self.il_style = self._capture_row_style(template_ws, il_row, max_cols)

        self.column_widths = {
            col_letter: template_ws.column_dimensions[col_letter].width
            for col_letter in template_ws.column_dimensions
        }

        self.workbook = template_wb
        self.worksheet = template_ws

        self._delete_other_sheets()
        self._clear_sheet_fully(self.worksheet)

        for col_letter, width in self.column_widths.items():
            if width:
                self.worksheet.column_dimensions[col_letter].width = width

        logging.info(
            f"Plantilla cargada dinámicamente. "
            f"HINV columnas: {len(self.hinv_headers)}, "
            f"HIL columnas: {len(self.hil_headers)}"
        )

    def _create_without_template(self):
        self.workbook = Workbook()
        self.worksheet = self.workbook.active
        self.worksheet.title = "Sheet1"

        self.hinv_headers = DEFAULT_HINV_HEADERS
        self.hil_headers = DEFAULT_HIL_HEADERS

        logging.info("Workbook creado sin plantilla. Se usan cabeceras por defecto.")

    def _find_marker_row(self, ws, marker):
        for row_idx in range(1, ws.max_row + 1):
            value = clean_cell(ws.cell(row=row_idx, column=1).value)

            if value == marker:
                return row_idx

        return None

    def _read_row_values(self, ws, row_idx):
        last_col = 0

        for col_idx in range(1, ws.max_column + 1):
            if clean_cell(ws.cell(row=row_idx, column=col_idx).value) is not None:
                last_col = col_idx

        if last_col == 0:
            return []

        return [
            clean_cell(ws.cell(row=row_idx, column=col_idx).value) or ""
            for col_idx in range(1, last_col + 1)
        ]

    def _capture_row_style(self, ws, row_idx, max_cols):
        row_height = ws.row_dimensions[row_idx].height
        styles = []

        for col_idx in range(1, max_cols + 1):
            cell = ws.cell(row=row_idx, column=col_idx)

            styles.append({
                "style": copy.copy(cell._style) if cell.has_style else None,
                "number_format": cell.number_format,
                "font": copy.copy(cell.font),
                "fill": copy.copy(cell.fill),
                "border": copy.copy(cell.border),
                "alignment": copy.copy(cell.alignment),
                "protection": copy.copy(cell.protection),
            })

        return {
            "height": row_height,
            "cells": styles
        }

    def _apply_row_style(self, row_idx, row_style, max_cols):
        if not row_style:
            return

        self.worksheet.row_dimensions[row_idx].height = row_style["height"]

        for col_idx in range(1, max_cols + 1):
            if col_idx > len(row_style["cells"]):
                break

            style_info = row_style["cells"][col_idx - 1]
            cell = self.worksheet.cell(row=row_idx, column=col_idx)

            if style_info["style"] is not None:
                cell._style = copy.copy(style_info["style"])

            cell.number_format = style_info["number_format"]
            cell.font = copy.copy(style_info["font"])
            cell.fill = copy.copy(style_info["fill"])
            cell.border = copy.copy(style_info["border"])
            cell.alignment = copy.copy(style_info["alignment"])
            cell.protection = copy.copy(style_info["protection"])

    def _delete_other_sheets(self):
        active = self.worksheet

        for ws in list(self.workbook.worksheets):
            if ws != active:
                self.workbook.remove(ws)

    def _clear_sheet_fully(self, ws):
        if ws.max_row > 0:
            ws.delete_rows(1, ws.max_row)

    def _write_row(self, row_idx, values, row_style):
        max_cols = len(values)

        self._apply_row_style(row_idx, row_style, max_cols)

        for col_idx, value in enumerate(values, start=1):
            cell = self.worksheet.cell(row=row_idx, column=col_idx)
            cell.value = value

            if isinstance(value, datetime):
                cell.number_format = self.DATE_FORMAT

    def _find_header_col(self, headers, target_header):
        target = normalize_header(target_header)

        for idx, header in enumerate(headers, start=1):
            if normalize_header(header) == target:
                return idx

        return None

    def _apply_line_type_dropdown(self, start_row, end_row):
        if end_row < start_row:
            return

        line_type_col = self._find_header_col(self.hil_headers, "Line type")

        if not line_type_col:
            return

        formula = '"' + ",".join(LINE_TYPE_OPTIONS) + '"'

        validation = DataValidation(
            type="list",
            formula1=formula,
            allow_blank=False
        )

        validation.error = "Selecciona un tipo de línea válido para Alma."
        validation.errorTitle = "Tipo de línea no válido"
        validation.prompt = "Selecciona un tipo de línea."
        validation.promptTitle = "Line type"

        self.worksheet.add_data_validation(validation)

        col_letter = get_column_letter(line_type_col)
        validation.add(f"{col_letter}{start_row}:{col_letter}{end_row}")

    def write_invoices(self, invoices):
        if self.workbook is None or self.worksheet is None:
            self.create_or_load_workbook()

        if len(invoices) != 1:
            raise ValueError(
                "Esta versión espera exactamente una factura por archivo para mantener "
                "la estructura Alma: HINV, INV, HIL, IL..."
            )

        invoice = invoices[0]

        row = 1

        self._write_row(row, self.hinv_headers, self.hinv_style)
        row += 1

        inv_row = self._build_invoice_row_from_headers(invoice, self.hinv_headers)
        self._write_row(row, inv_row, self.inv_style)
        row += 1

        self._write_row(row, self.hil_headers, self.hil_style)
        row += 1

        first_il_row = row

        for line in invoice.lines:
            il_row = self._build_line_row_from_headers(line, self.hil_headers)
            self._write_row(row, il_row, self.il_style)
            row += 1

        last_il_row = row - 1

        self._apply_line_type_dropdown(first_il_row, last_il_row)

    def _build_invoice_row_from_headers(self, invoice, headers):
        values = {
            "HINV": "INV",
            "Invoice Number": invoice.invoice_number,
            "Date": invoice.date,
            "Total amount": float(invoice.total_amount),
            "Currency": invoice.currency,
            "Vendor Account": invoice.vendor_account,
            "Payment method": None,
            "Note": invoice.note,
            "Use Pro Rata": None,
            "Shipment amount": None,
            "Overhead amount": None,
            "Insurance amount": None,
            "Discount amount": None,
            "VAT Percentage": None,
            "VAT Amount": decimal_to_float_or_none(invoice.vat_amount),
            "Inclusive": invoice.inclusive,
            "Expended From Fund": None,
            "PrePaid": None,
            "Payment status": None,
            "Voucher number": None,
            "Voucher date": None,
            "Voucher Amount": None,
            "Invoice Reference Number": None,
            "VAT In Invoice Line Level": "YES",
            "Report TAX": None,
            "VAT code": None,
        }

        return [
            values.get(normalize_header(header), None)
            for header in headers
        ]

    def _build_line_row_from_headers(self, line, headers):
        values = {
            "HIL": "IL",
            "Line type": "REGULAR",
            "Line Number": line.line_number,
            "PO Line": line.po_line,
            "Vendor Ref Number": None,
            "ISSN": line.issn,
            "ISBN": line.isbn,
            "Title": line.title,
            "Price": float(line.price),
            "Price Note": None,
            "Quantity": float(line.quantity),
            "Reporting Code": line.reporting_code,
            "Secondary Reporting Code": line.secondary_reporting_code,
            "Tertiary Reporting Code": line.tertiary_reporting_code,
            "Fourth Reporting Code": line.fourth_reporting_code,
            "Fifth Reporting Code": line.fifth_reporting_code,
            "VAT code": None,
            "Start subs date": None,
            "End subs date": None,
            "Additional Info": None,
            "Note": None,
            "VAT Amount": decimal_to_float_or_none(line.vat_amount),
            "VAT Percentage": decimal_to_float_or_none(line.vat_percentage),
        }

        row = []
        fund_index = 0

        for header in headers:
            normalized_header = normalize_header(header)

            if normalized_header == "Fund and percent":
                if fund_index < len(line.funds):
                    row.append(line.funds[fund_index])
                else:
                    row.append(None)

                fund_index += 1
            else:
                row.append(values.get(normalized_header, None))

        return row

    def save(self):
        try:
            self.workbook.save(self.output_path)
            logging.info(f"Excel Alma guardado correctamente: {self.output_path}")

        except PermissionError:
            logging.exception("No se puede guardar el Excel porque está abierto o bloqueado")
            raise ValueError(
                f"No se pudo guardar el archivo '{self.output_path}'. "
                "Puede que esté abierto en Excel."
            )

        except Exception as e:
            logging.exception("Error guardando Excel")
            raise ValueError(f"Error al guardar el Excel: {e}")


# ============================================================
# EXPORTADOR DE VALIDACIÓN
# ============================================================

class ValidationReportExporter:
    def __init__(self, output_path):
        self.output_path = output_path

    def save(self, issues):
        if not issues:
            return None

        wb = Workbook()
        ws = wb.active
        ws.title = "Validación"

        ws.append([
            "Tipo aviso",
            "Línea factura",
            "ISBN XSIG/XML",
            "ISSN XSIG/XML",
            "Título",
            "PO Lines candidatas",
            "Mensaje"
        ])

        for issue in issues:
            ws.append([
                issue.issue_type,
                issue.invoice_line_number,
                issue.isbn,
                issue.issn,
                issue.title,
                "; ".join(issue.candidates),
                issue.message
            ])

        wb.save(self.output_path)
        logging.info(f"Informe de validación guardado: {self.output_path}")
        return self.output_path


# ============================================================
# FUNCIÓN PRINCIPAL
# ============================================================

def convert_facturae_to_alma_excel(
    invoice_path,
    alma_pol_excel_path,
    excel_output_path,
    template_path=None
):
    parser = FacturaeParser(invoice_path)
    invoices = parser.parse()

    if len(invoices) != 1:
        raise ValueError(
            f"El archivo contiene {len(invoices)} facturas. "
            "Esta versión espera una única factura por archivo."
        )

    lookup = AlmaPOLineLookup(alma_pol_excel_path)
    lookup.load()

    enricher = AlmaLineEnricher(lookup)
    validation_issues = enricher.enrich(invoices)

    exporter = AlmaExcelExporter(
        output_path=excel_output_path,
        template_path=template_path
    )

    exporter.create_or_load_workbook()
    exporter.write_invoices(invoices)
    exporter.save()

    validation_path = None

    if validation_issues:
        validation_path = build_validation_output_path(excel_output_path)
        ValidationReportExporter(validation_path).save(validation_issues)

    summary = build_conversion_summary(invoices, validation_issues)

    return len(invoices), len(validation_issues), validation_path, summary


# ============================================================
# INTERFAZ TKINTER
# ============================================================

class FacturaeApp:
    def __init__(self, root):
        self.root = root
        self.root.title(APP_TITLE)
        self.root.geometry("980x640")
        self.root.resizable(False, False)
        self.root.configure(bg=UI_BG)

        self.invoice_file = tk.StringVar()
        self.alma_pol_file = tk.StringVar()
        self.template_file = tk.StringVar()
        self.output_file = tk.StringVar()

        self.status_text = tk.StringVar(
            value="Estado: pendiente de seleccionar archivos."
        )

        self._setup_style()
        self._build_ui()

    def _setup_style(self):
        style = ttk.Style()

        try:
            style.theme_use("clam")
        except Exception:
            pass

        style.configure("App.TFrame", background=UI_BG)
        style.configure("App.TLabel", background=UI_BG)
        style.configure("App.TLabelframe", background=UI_BG)
        style.configure("App.TLabelframe.Label", background=UI_BG)

    def _open_url(self, url):
        try:
            webbrowser.open_new(url)
        except Exception as e:
            messagebox.showerror(
                "Error al abrir enlace",
                f"No se pudo abrir el enlace:\n{url}\n\n{e}"
            )

    def _open_folder(self, folder_path):
        if not folder_path or not os.path.exists(folder_path):
            messagebox.showwarning(
                "Aviso",
                "No se encontró la carpeta de salida."
            )
            return

        try:
            system_name = platform.system()

            if system_name == "Windows":
                os.startfile(folder_path)
            elif system_name == "Darwin":
                subprocess.Popen(["open", folder_path])
            else:
                subprocess.Popen(["xdg-open", folder_path])

        except Exception as e:
            messagebox.showerror(
                "Error al abrir carpeta",
                f"No se pudo abrir la carpeta:\n{folder_path}\n\n{e}"
            )

    def open_output_folder(self):
        output_path = self.output_file.get().strip()

        if not output_path:
            messagebox.showwarning(
                "Aviso",
                "No hay archivo de salida definido."
            )
            return

        folder = os.path.dirname(output_path)

        if not folder:
            messagebox.showwarning(
                "Aviso",
                "No se pudo determinar la carpeta de salida."
            )
            return

        self._open_folder(folder)

    def _build_ui(self):
        main = ttk.Frame(self.root, padding=20, style="App.TFrame")
        main.pack(fill=tk.BOTH, expand=True)
        main.columnconfigure(0, weight=1)

        title = ttk.Label(
            main,
            text=APP_TITLE,
            font=("Segoe UI", 14, "bold"),
            style="App.TLabel"
        )
        title.grid(row=0, column=0, sticky="w", pady=(0, 15))

        input_frame = ttk.LabelFrame(
            main,
            text="1. Archivos de entrada",
            padding=12,
            style="App.TLabelframe"
        )
        input_frame.grid(row=1, column=0, sticky="ew")
        input_frame.columnconfigure(0, weight=1)

        ttk.Label(
            input_frame,
            text="Factura electrónica XSIG / XML / TXT *",
            style="App.TLabel"
        ).grid(row=0, column=0, sticky="w")

        ttk.Entry(
            input_frame,
            textvariable=self.invoice_file
        ).grid(row=1, column=0, sticky="ew", padx=(0, 10), pady=(2, 10))

        ttk.Button(
            input_frame,
            text="Seleccionar factura...",
            command=self.select_invoice
        ).grid(row=1, column=1, sticky="ew", pady=(2, 10))

        ttk.Label(
            input_frame,
            text="Excel de Alma con PO Lines, Title, ISBN e ISSN *",
            style="App.TLabel"
        ).grid(row=2, column=0, sticky="w")

        ttk.Entry(
            input_frame,
            textvariable=self.alma_pol_file
        ).grid(row=3, column=0, sticky="ew", padx=(0, 10), pady=(2, 10))

        ttk.Button(
            input_frame,
            text="Seleccionar Excel Alma...",
            command=self.select_alma_pol
        ).grid(row=3, column=1, sticky="ew", pady=(2, 10))

        ttk.Label(
            input_frame,
            text="Plantilla Alma .xlsx para leer columnas/formato *",
            style="App.TLabel"
        ).grid(row=4, column=0, sticky="w")

        ttk.Entry(
            input_frame,
            textvariable=self.template_file
        ).grid(row=5, column=0, sticky="ew", padx=(0, 10), pady=(2, 0))

        ttk.Button(
            input_frame,
            text="Seleccionar plantilla...",
            command=self.select_template
        ).grid(row=5, column=1, sticky="ew", pady=(2, 0))

        output_frame = ttk.LabelFrame(
            main,
            text="2. Archivo de salida",
            padding=12,
            style="App.TLabelframe"
        )
        output_frame.grid(row=2, column=0, sticky="ew", pady=(12, 0))
        output_frame.columnconfigure(0, weight=1)

        ttk.Label(
            output_frame,
            text="Guardar Excel final como *",
            style="App.TLabel"
        ).grid(row=0, column=0, sticky="w")

        ttk.Entry(
            output_frame,
            textvariable=self.output_file
        ).grid(row=1, column=0, sticky="ew", padx=(0, 10), pady=(2, 0))

        ttk.Button(
            output_frame,
            text="Elegir destino...",
            command=self.select_output
        ).grid(row=1, column=1, sticky="ew", pady=(2, 0))

        actions_frame = ttk.LabelFrame(
            main,
            text="3. Acciones",
            padding=12,
            style="App.TLabelframe"
        )
        actions_frame.grid(row=3, column=0, sticky="ew", pady=(12, 0))

        for col in range(4):
            actions_frame.columnconfigure(col, weight=1)

        self.convert_button = ttk.Button(
            actions_frame,
            text="Convertir a formato Alma",
            command=self.run_conversion
        )
        self.convert_button.grid(row=0, column=0, sticky="ew", padx=(0, 5))

        ttk.Button(
            actions_frame,
            text="Limpiar datos",
            command=self.clear_fields
        ).grid(row=0, column=1, sticky="ew", padx=5)

        ttk.Button(
            actions_frame,
            text="Abrir carpeta de salida",
            command=self.open_output_folder
        ).grid(row=0, column=2, sticky="ew", padx=5)

        ttk.Button(
            actions_frame,
            text="Salir",
            command=self.root.destroy
        ).grid(row=0, column=3, sticky="ew", padx=(5, 0))

        status_label = ttk.Label(
            main,
            textvariable=self.status_text,
            style="App.TLabel",
            font=("Segoe UI", 9, "italic")
        )
        status_label.grid(row=4, column=0, sticky="w", pady=(12, 0))

        info = (
            "Notas:\n"
            "- XSIG es la opción principal, aunque también se admiten XML y TXT.\n"
            "- La plantilla Alma se usa para leer el número y orden exacto de columnas.\n"
            "- PO Line solo se toma del Excel de Alma, nunca del XSIG.\n"
            "- Si no hay ISBN/ISSN, se intenta match por título usando stopwords integradas.\n"
            "- Si no hay coincidencia, se deja PO Line vacío y se genera validación.\n"
            "- Los campos marcados con * son obligatorios."
        )

        ttk.Label(
            main,
            text=info,
            justify="left",
            style="App.TLabel"
        ).grid(row=5, column=0, sticky="w", pady=(12, 0))

        repo_label = tk.Label(
            main,
            text=f"Repositorio del proyecto y descarga: {PROJECT_REPOSITORY_URL}",
            fg="blue",
            bg=UI_BG,
            cursor="hand2",
            font=("Segoe UI", 9, "underline")
        )
        repo_label.grid(row=6, column=0, sticky="w", pady=(8, 0))
        repo_label.bind(
            "<Button-1>",
            lambda event: self._open_url(PROJECT_REPOSITORY_URL)
        )

        guide_label = tk.Label(
            main,
            text=f"Instrucciones en Biblioguía Facturae to Alma: {PROJECT_GUIDE_URL}",
            fg="blue",
            bg=UI_BG,
            cursor="hand2",
            font=("Segoe UI", 9, "underline")
        )
        guide_label.grid(row=7, column=0, sticky="w", pady=(2, 0))
        guide_label.bind(
            "<Button-1>",
            lambda event: self._open_url(PROJECT_GUIDE_URL)
        )

    def select_invoice(self):
        path = filedialog.askopenfilename(
            title="Seleccionar factura XSIG / XML / TXT",
            filetypes=[
                ("Facturas Universitas XXI / Facturae", "*.xsig *.xml *.txt"),
                ("Archivos XSIG", "*.xsig"),
                ("Archivos XML", "*.xml"),
                ("Archivos TXT", "*.txt"),
                ("Todos los archivos", "*.*")
            ]
        )

        if path:
            self.invoice_file.set(path)

            base, _ = os.path.splitext(path)
            self.output_file.set(base + "_Alma.xlsx")

            self.status_text.set(
                "Estado: factura seleccionada. Revisa el Excel de Alma, la plantilla y el archivo de salida."
            )

    def select_alma_pol(self):
        path = filedialog.askopenfilename(
            title="Seleccionar Excel de Alma con PO Lines",
            filetypes=[
                ("Archivos Excel", "*.xlsx"),
                ("Todos los archivos", "*.*")
            ]
        )

        if path:
            self.alma_pol_file.set(path)
            self.status_text.set("Estado: Excel de Alma seleccionado.")

    def select_template(self):
        path = filedialog.askopenfilename(
            title="Seleccionar plantilla Alma .xlsx",
            filetypes=[
                ("Archivos Excel", "*.xlsx"),
                ("Todos los archivos", "*.*")
            ]
        )

        if path:
            self.template_file.set(path)
            self.status_text.set("Estado: plantilla Alma seleccionada.")

    def select_output(self):
        path = filedialog.asksaveasfilename(
            title="Definir archivo Excel de salida",
            defaultextension=".xlsx",
            filetypes=[
                ("Archivos Excel", "*.xlsx"),
                ("Todos los archivos", "*.*")
            ]
        )

        if path:
            self.output_file.set(path)
            self.status_text.set("Estado: archivo de salida definido.")

    def clear_fields(self):
        self.invoice_file.set("")
        self.output_file.set("")

        if self.alma_pol_file.get().strip():
            clear_alma = messagebox.askyesno(
                "Limpiar Excel de Alma",
                "¿Quieres limpiar también el Excel de Alma con PO Lines?\n\n"
                "La plantilla Alma se conservará."
            )

            if clear_alma:
                self.alma_pol_file.set("")

        self.status_text.set("Estado: datos limpiados. La plantilla Alma se conserva.")

    def _build_summary_message(self, summary):
        if not summary:
            return ""

        return (
            "\n\nResumen de líneas:\n"
            f"- Líneas procesadas: {summary['total_lines']}\n"
            f"- Coincidencias por ISBN: {summary['match_isbn']}\n"
            f"- Coincidencias por ISSN: {summary['match_issn']}\n"
            f"- Coincidencias por título: {summary['match_title']}\n"
            f"- Duplicados ISBN: {summary['duplicate_isbn']}\n"
            f"- Duplicados ISSN: {summary['duplicate_issn']}\n"
            f"- Títulos ambiguos: {summary['ambiguous_title']}\n"
            f"- Sin coincidencia: {summary['no_match']}\n"
            f"- Líneas sin PO Line: {summary['lines_without_po_line']}"
        )

    def run_conversion(self):
        invoice_path = self.invoice_file.get().strip()
        alma_pol_path = self.alma_pol_file.get().strip()
        template_path = self.template_file.get().strip() or None
        output_path = self.output_file.get().strip()

        if not invoice_path:
            messagebox.showerror("Error", "Debes seleccionar el archivo de factura XSIG, XML o TXT.")
            self.status_text.set("Estado: falta seleccionar la factura.")
            return

        allowed_extensions = {".xsig", ".xml", ".txt"}
        ext = os.path.splitext(invoice_path)[1].lower()

        if ext not in allowed_extensions:
            messagebox.showerror(
                "Error",
                "El archivo de factura debe tener extensión .xsig, .xml o .txt."
            )
            self.status_text.set("Estado: extensión de factura no válida.")
            return

        if not alma_pol_path:
            messagebox.showerror("Error", "Debes seleccionar el Excel de Alma con PO Lines.")
            self.status_text.set("Estado: falta seleccionar el Excel de Alma.")
            return

        if not template_path:
            messagebox.showerror("Error", "Debes seleccionar la plantilla Alma .xlsx.")
            self.status_text.set("Estado: falta seleccionar la plantilla Alma.")
            return

        if not output_path:
            messagebox.showerror("Error", "Debes definir el archivo Excel de salida.")
            self.status_text.set("Estado: falta definir el archivo de salida.")
            return

        if not os.path.exists(invoice_path):
            messagebox.showerror("Error", "El archivo de factura seleccionado no existe.")
            self.status_text.set("Estado: el archivo de factura no existe.")
            return

        if not os.path.exists(alma_pol_path):
            messagebox.showerror("Error", "El Excel de Alma seleccionado no existe.")
            self.status_text.set("Estado: el Excel de Alma no existe.")
            return

        if template_path and not os.path.exists(template_path):
            messagebox.showerror("Error", "La plantilla Alma seleccionada no existe.")
            self.status_text.set("Estado: la plantilla Alma no existe.")
            return

        self.convert_button.config(state="disabled")
        self.status_text.set("Estado: convirtiendo factura...")
        self.root.update_idletasks()

        try:
            total_invoices, total_issues, validation_path, summary = convert_facturae_to_alma_excel(
                invoice_path=invoice_path,
                alma_pol_excel_path=alma_pol_path,
                excel_output_path=output_path,
                template_path=template_path
            )

            message = (
                "Conversión realizada correctamente.\n\n"
                f"Facturas procesadas: {total_invoices}\n"
                f"Avisos de validación: {total_issues}\n\n"
                f"Excel Alma generado:\n{output_path}"
            )

            message += self._build_summary_message(summary)

            if validation_path:
                message += (
                    "\n\nTambién se ha generado un informe de validación:\n"
                    f"{validation_path}"
                )

            if summary and summary.get("lines_without_po_line", 0) > 0:
                message += (
                    "\n\nAtención: hay líneas sin PO Line. "
                    "Revisa el fichero de validación antes de cargar en Alma."
                )

            self.status_text.set("Estado: conversión completada.")

            messagebox.showinfo("Conversión completada", message)

        except Exception as e:
            logging.exception("Error durante la conversión")
            traceback.print_exc()

            self.status_text.set("Estado: error durante la conversión.")

            messagebox.showerror(
                "Error durante la conversión",
                f"{str(e)}\n\n"
                f"Consulta el archivo de log '{LOG_FILE}' para más detalles."
            )

        finally:
            self.convert_button.config(state="normal")


# ============================================================
# MAIN
# ============================================================

if __name__ == "__main__":
    root = tk.Tk()
    app = FacturaeApp(root)
    root.mainloop()
