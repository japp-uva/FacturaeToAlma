# SPDX-License-Identifier: GPL-3.0-only
# FacturaeToAlma 3.1

import os
import sys
import json
import queue
import threading
import re
import copy
import html
import logging
import traceback
import unicodedata
import webbrowser
import subprocess  # nosec B404 - uso revisado, sin shell y con ejecutable absoluto
import platform
import shutil
from dataclasses import dataclass, field
from datetime import datetime, date
from decimal import Decimal, InvalidOperation
from difflib import SequenceMatcher

import tkinter as tk
from tkinter import ttk, filedialog, messagebox

from PIL import Image, ImageTk, UnidentifiedImageError

import defusedxml.ElementTree as ET
from defusedxml.common import DefusedXmlException

from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

APP_NAME = "FacturaeToAlma"
APP_VERSION = "3.1"
APP_TITLE = f"Universidad de Valladolid. Biblioteca -> {APP_NAME} {APP_VERSION}"
LOG_FILE = "facturae_alma.log"
PROJECT_REPOSITORY_URL = "https://github.com/japp-uva/FacturaeToAlma"
PROJECT_GUIDE_URL = "https://biblioguias.uva.es/facturae_to_alma"
UI_BG = "#f0f0f0"
TITLE_MATCH_MIN_SCORE = 0.78
TITLE_MATCH_MIN_MARGIN = 0.08
LINE_TYPE_OPTIONS = ["REGULAR", "OVERHEAD", "OTHER", "SHIPMENT", "DISCOUNT", "INSURANCE", "ADDITIONAL_CHARGES"]
DANGEROUS_EXCEL_PREFIXES = ("=", "+", "-", "@", "\t", "\r", "\n")
MAX_LOGO_HEIGHT = 100

logging.basicConfig(filename=LOG_FILE, level=logging.INFO, format="%(asctime)s | %(levelname)s | %(message)s")

VENDOR_MAP = {"DEFAULT": "PROV_GENERICO"}

DEFAULT_HINV_HEADERS = [
    "HINV", "Invoice Number", "Date", "Total amount", "Currency", "Vendor Account", "Payment method", "Note",
    "Use Pro Rata", "Shipment amount", "Overhead amount", "Insurance amount", "Discount amount", "VAT Percentage",
    "VAT Amount", "Inclusive", "Expended From Fund", "PrePaid", "Payment status", "Voucher number", "Voucher date",
    "Voucher Amount", "Invoice Reference Number", "VAT In Invoice Line Level", "Report TAX"
]
DEFAULT_HIL_HEADERS = [
    "HIL", "Line type", "Line Number", "PO Line", "Vendor Ref Number", "ISSN", "ISBN", "Title", "Price", "Quantity",
    "Reporting Code", "Secondary Reporting Code", "Tertiary Reporting Code", "Fourth Reporting Code", "Fifth Reporting Code",
    "Start subs date", "End subs date", "Note", "Fund and percent", "Fund and percent", "VAT Amount", "VAT Percentage"
]

STOPWORDS = {
    "el", "la", "los", "las", "un", "una", "unos", "unas", "a", "ante", "bajo", "con", "contra", "de", "desde", "en", "entre", "hacia", "hasta", "para", "por", "segun", "según", "sin", "sobre", "tras", "y", "e", "ni", "que", "pero", "aunque", "sino", "porque", "si", "sí", "me", "te", "se", "nos", "mi", "mí", "tu", "tú", "su", "este", "esta", "esto",
    "an", "the", "in", "on", "at", "by", "for", "with", "about", "from", "to", "of", "under", "over", "and", "but", "or", "so", "because", "although", "yet", "i", "you", "he", "she", "it", "we", "they", "this", "that", "there",
    "le", "les", "l", "une", "des", "à", "dans", "sur", "avec", "pour", "par", "sans", "chez", "et", "ou", "mais", "donc", "car", "je", "il", "elle", "nous", "vous", "ce",
    "der", "die", "das", "dem", "den", "ein", "eine", "einer", "auf", "zu", "mit", "von", "fur", "für", "bei", "durch", "aus", "vor", "und", "oder", "aber", "denn", "weil", "obwohl", "sondern", "ich", "du", "er", "sie", "es", "wir", "ihr", "sich",
    "il", "lo", "gli", "di", "da", "su", "per", "tra", "fra", "ed", "ma", "oppure", "perche", "perché", "anche", "ti", "ci", "io", "lui", "lei", "noi", "voi",
    "els", "amb", "sense", "sota", "perque", "perquè", "doncs", "jo", "ell", "ella", "nosaltres", "vosaltres", "em", "et",
    "o", "os", "as", "unha", "baixo", "segundo", "sen", "senon", "senón", "mais", "nin", "eu", "ela", "nós", "vós", "vos",
    "ak", "ari", "ariak", "bat", "batzuk", "baino", "gabe", "gora", "behera", "arte", "bila", "kontra", "zehar", "aurka", "eta", "edo", "baina", "denez", "beraz", "orduan", "baita", "ezta", "zein", "hi", "hura", "gu", "zu", "zuek", "haiek", "hau", "hori",
    "um", "uma", "umas", "com", "em", "sem", "sob", "trás", "ate", "até", "mas", "nem", "embora", "contudo", "pois", "ele", "isto", "esse", "essa", "isso"
}
STOPWORD_PHRASES = {"per a", "des de", "cap a", "fins a", "encara que", "ainda que", "aínda que"}

@dataclass
class AlmaPOLRecord:
    po_line: str
    title: str = None
    reporting_code: str = None
    secondary_reporting_code: str = None
    tertiary_reporting_code: str = None
    fourth_reporting_code: str = None
    fifth_reporting_code: str = None
    start_subs_date: datetime = None
    isbn_raw: str = None
    issn_raw: str = None
    quantity_for_pricing: Decimal = None
    funds: list = field(default_factory=list)
    normalized_title_words: list = field(default_factory=list)
    normalized_title_text: str = ""

@dataclass
class ValidationIssue:
    issue_type: str
    invoice_number: str
    invoice_line_number: int
    isbn: str
    issn: str
    title: str
    candidates: list
    message: str
    po_line: str = None
    invoice_quantity: Decimal = None
    alma_quantity_for_pricing: Decimal = None

@dataclass
class InvoiceLineData:
    line_number: int
    po_line: str = None
    vendor_ref: str = None
    issn: str = None
    isbn: str = None
    title: str = None
    price: Decimal = Decimal("0.00")
    quantity: Decimal = Decimal("1.00")
    start_date: datetime = None
    end_date: datetime = None
    note: str = None
    vat_amount: Decimal = None
    vat_percentage: Decimal = None
    reporting_code: str = None
    secondary_reporting_code: str = None
    tertiary_reporting_code: str = None
    fourth_reporting_code: str = None
    fifth_reporting_code: str = None
    alma_quantity_for_pricing: Decimal = None
    funds: list = field(default_factory=list)
    match_status: str = None
    match_message: str = None
    match_candidates: list = field(default_factory=list)

@dataclass
class InvoiceData:
    invoice_number: str
    date: datetime
    total_amount: Decimal
    currency: str
    vendor_account: str
    note: str
    seller_name: str = None
    seller_tax_id: str = None
    source_file: str = None
    vat_amount: Decimal = None
    inclusive: bool = False
    lines: list = field(default_factory=list)

def sanitize_excel_value(value):
    if isinstance(value, str) and value.startswith(DANGEROUS_EXCEL_PREFIXES):
        return "'" + value
    return value

def clean_cell(value):
    if value is None:
        return None
    text = str(value).replace("\xa0", " ").strip()
    if not text or text.lower() == "nan":
        return None
    return text

def alma_empty_to_none(value):
    value = clean_cell(value)
    return None if value in (None, "", "-1") else value

def safe_decimal(value, default=Decimal("0.00")):
    if value is None:
        return default
    raw = str(value).strip().replace(" ", "")
    if not raw:
        return default
    if "." in raw and "," in raw:
        if raw.rfind(",") > raw.rfind("."):
            raw = raw.replace(".", "").replace(",", ".")
        else:
            raw = raw.replace(",", "")
    elif "," in raw and "." not in raw:
        raw = raw.replace(",", ".")
    try:
        return Decimal(raw)
    except InvalidOperation:
        logging.warning(f"No se pudo convertir a Decimal: {value!r}")
        return default

def decimal_or_none(value):
    if value is None or not str(value).strip():
        return None
    return safe_decimal(value)

def decimal_to_float_or_none(value):
    return None if value is None else float(value)

def parse_date_for_excel(value):
    if value is None:
        return None
    if isinstance(value, datetime):
        return value
    if isinstance(value, date):
        return datetime(value.year, value.month, value.day)
    text = clean_cell(value)
    if not text:
        return None
    text = text[:10]
    for fmt in ["%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y", "%m/%d/%Y", "%Y/%m/%d"]:
        try:
            return datetime.strptime(text, fmt)
        except ValueError:
            pass
    logging.warning(f"No se pudo convertir la fecha: {value!r}")
    return None

def normalize_identifier(value):
    if not value:
        return ""
    return re.sub(r"[^0-9X]", "", str(value).replace("\xa0", " ").upper())

def split_identifiers(value):
    value = clean_cell(value)
    if not value:
        return []
    return [norm for norm in (normalize_identifier(part) for part in value.split(";")) if norm]

def is_valid_isbn10(code):
    code = normalize_identifier(code)
    if not re.fullmatch(r"\d{9}[\dX]", code):
        return False
    total = sum((10 - i) * (10 if c == "X" else int(c)) for i, c in enumerate(code))
    return total % 11 == 0

def is_valid_isbn13(code):
    code = normalize_identifier(code)
    if not re.fullmatch(r"\d{13}", code):
        return False
    total = sum(int(c) if i % 2 == 0 else int(c) * 3 for i, c in enumerate(code[:12]))
    check = (10 - (total % 10)) % 10
    return check == int(code[-1])

def is_valid_issn(code):
    code = normalize_identifier(code)
    if not re.fullmatch(r"\d{7}[\dX]", code):
        return False
    total = sum(int(code[i]) * (8 - i) for i in range(7))
    check = (11 - (total % 11)) % 11
    expected = "X" if check == 10 else str(check)
    return code[-1] == expected

def looks_like_issn(code):
    return bool(re.fullmatch(r"\d{7}[\dX]", normalize_identifier(code)))

def looks_like_isbn(code):
    return bool(re.fullmatch(r"\d{10}|\d{13}", normalize_identifier(code)))

def extract_identifiers_from_text(text):
    if not text:
        return None, None, text
    original_text = str(text)
    isbn = None
    issn = None
    matched = []
    candidate_pattern = re.compile(r"[0-9Xx][0-9Xx\-\s]{6,25}[0-9Xx]")
    for match in candidate_pattern.finditer(original_text):
        fragment = match.group(0)
        normalized = normalize_identifier(fragment)
        if not isbn and len(normalized) == 13 and is_valid_isbn13(normalized):
            isbn = normalized
            matched.append(fragment)
            continue
        if not isbn and len(normalized) == 10 and is_valid_isbn10(normalized):
            isbn = normalized
            matched.append(fragment)
            continue
        if not issn and len(normalized) == 8 and is_valid_issn(normalized):
            issn = normalized
            matched.append(fragment)
            continue
    cleaned = original_text
    for fragment in matched:
        cleaned = cleaned.replace(fragment, " ")
    cleaned = re.sub(r"\s*/\s*$", "", cleaned)
    cleaned = re.sub(r"\s*[-–—:;,/]\s*$", "", cleaned)
    cleaned = re.sub(r"\s+", " ", cleaned).strip()
    return isbn, issn, cleaned

def classify_article_code(code):
    if not code:
        return None, None, None
    isbn, issn, _ = extract_identifiers_from_text(code)
    if isbn:
        return None, isbn, None
    if issn:
        return issn, None, None
    if looks_like_issn(code):
        return code, None, None
    if looks_like_isbn(code):
        return None, code, None
    return None, None, code

def normalize_header(header):
    return "" if header is None else " ".join(str(header).strip().split())

def get_vendor_account(nif, seller_name=None):
    nif = nif.strip() if nif else ""
    if nif and nif in VENDOR_MAP:
        return VENDOR_MAP[nif]
    if seller_name:
        return seller_name.strip()
    return VENDOR_MAP["DEFAULT"]

def build_validation_output_path(output_path):
    base, _ = os.path.splitext(output_path)
    return f"{base}_validacion.xlsx"

def remove_accents(text):
    if not text:
        return ""
    normalized = unicodedata.normalize("NFD", text)
    return "".join(char for char in normalized if unicodedata.category(char) != "Mn")

def normalize_title_for_matching(title):
    if not title:
        return []
    text = remove_accents(html.unescape(str(title)).lower()).replace("’", "'").replace("´", "'")
    for phrase in STOPWORD_PHRASES:
        normalized_phrase = re.sub(r"[^a-z0-9]+", " ", remove_accents(phrase.lower())).strip()
        if normalized_phrase:
            text = re.sub(rf"\b{re.escape(normalized_phrase)}\b", " ", text)
    text = re.sub(r"[^a-z0-9]+", " ", text)
    return [w for w in text.split() if len(w) > 1 and w not in STOPWORDS]

def title_similarity_score(invoice_title, record):
    invoice_words = normalize_title_for_matching(invoice_title)
    alma_words = record.normalized_title_words
    if not invoice_words or not alma_words:
        return 0.0
    invoice_set = set(invoice_words)
    alma_set = set(alma_words)
    common = invoice_set.intersection(alma_set)
    word_score = len(common) / len(invoice_set)
    sequence_score = SequenceMatcher(None, " ".join(invoice_words), record.normalized_title_text).ratio()
    return word_score * 0.75 + sequence_score * 0.25

def normalize_provider_for_comparison(value):
    if not value:
        return ""
    text = remove_accents(str(value).lower())
    text = re.sub(r"[^a-z0-9]+", " ", text)
    return re.sub(r"\s+", " ", text).strip()

def get_vendor_batch_key(invoice):
    if invoice.seller_tax_id:
        return normalize_identifier(invoice.seller_tax_id)
    if invoice.vendor_account:
        return normalize_provider_for_comparison(invoice.vendor_account)
    return ""

def build_conversion_summary(invoices, validation_issues):
    summary = {"total_lines": 0, "match_isbn": 0, "match_issn": 0, "match_title": 0, "duplicate_isbn": 0, "duplicate_issn": 0, "duplicate_title": 0, "ambiguous_title": 0, "no_match": 0, "validation_issues": len(validation_issues), "lines_without_po_line": 0}
    for invoice in invoices:
        for line in invoice.lines:
            summary["total_lines"] += 1
            status = line.match_status or ""
            if status == "MATCH_ISBN": summary["match_isbn"] += 1
            elif status == "MATCH_ISSN": summary["match_issn"] += 1
            elif status == "MATCH_TITLE": summary["match_title"] += 1
            elif status == "DUPLICATE_ISBN": summary["duplicate_isbn"] += 1
            elif status == "DUPLICATE_ISSN": summary["duplicate_issn"] += 1
            elif status == "DUPLICATE_TITLE": summary["duplicate_title"] += 1
            elif status == "AMBIGUOUS_TITLE": summary["ambiguous_title"] += 1
            elif status == "NO_MATCH": summary["no_match"] += 1
            if not line.po_line:
                summary["lines_without_po_line"] += 1
    return summary

class AlmaPOLineLookup:
    REQUIRED_COLUMNS = ["PO Line", "Reporting Code", "Secondary Reporting Code", "Tertiary Reporting Code", "Start subs date", "Title", "ISSN", "ISBN", "Quantity for Pricing"]
    def __init__(self, alma_excel_path):
        self.alma_excel_path = alma_excel_path
        self.records = []
        self.isbn_index = {}
        self.issn_index = {}
    def load(self):
        try:
            wb = load_workbook(self.alma_excel_path, data_only=True)
            ws = wb.active
        except Exception as e:
            logging.exception("Error cargando Excel de Alma")
            raise ValueError(f"No se pudo leer el Excel de Alma con PO Lines: {e}")
        headers = self._read_headers(ws)
        self._validate_headers(headers)
        col_map = {}
        col_map_multi = {}
        for idx, header in enumerate(headers, start=1):
            normalized = normalize_header(header)
            if normalized not in col_map:
                col_map[normalized] = idx
            col_map_multi.setdefault(normalized, []).append(idx)
        fund_columns = col_map_multi.get("Fund and percent", [])
        for row_idx in range(2, ws.max_row + 1):
            po_line = clean_cell(ws.cell(row=row_idx, column=col_map["PO Line"]).value)
            if not po_line:
                continue
            title = clean_cell(ws.cell(row=row_idx, column=col_map["Title"]).value)
            title_words = normalize_title_for_matching(title)
            funds = []
            for fund_col in fund_columns:
                fund_value = alma_empty_to_none(ws.cell(row=row_idx, column=fund_col).value)
                if fund_value:
                    funds.append(fund_value)
            rec = AlmaPOLRecord(
                po_line=po_line,
                reporting_code=alma_empty_to_none(ws.cell(row=row_idx, column=col_map["Reporting Code"]).value),
                secondary_reporting_code=alma_empty_to_none(ws.cell(row=row_idx, column=col_map["Secondary Reporting Code"]).value),
                tertiary_reporting_code=alma_empty_to_none(ws.cell(row=row_idx, column=col_map["Tertiary Reporting Code"]).value),
                fourth_reporting_code=alma_empty_to_none(ws.cell(row=row_idx, column=col_map["Fourth Reporting Code"]).value) if "Fourth Reporting Code" in col_map else None,
                fifth_reporting_code=alma_empty_to_none(ws.cell(row=row_idx, column=col_map["Fifth Reporting Code"]).value) if "Fifth Reporting Code" in col_map else None,
                start_subs_date=parse_date_for_excel(ws.cell(row=row_idx, column=col_map["Start subs date"]).value),
                title=title,
                issn_raw=clean_cell(ws.cell(row=row_idx, column=col_map["ISSN"]).value),
                quantity_for_pricing=decimal_or_none(ws.cell(row=row_idx, column=col_map["Quantity for Pricing"]).value),
                isbn_raw=clean_cell(ws.cell(row=row_idx, column=col_map["ISBN"]).value),
                funds=funds,
                normalized_title_words=title_words,
                normalized_title_text=" ".join(title_words)
            )
            self.records.append(rec)
            for isbn in split_identifiers(rec.isbn_raw):
                self.isbn_index.setdefault(isbn, []).append(rec)
            for issn in split_identifiers(rec.issn_raw):
                self.issn_index.setdefault(issn, []).append(rec)
        logging.info(f"Excel Alma cargado: {len(self.records)} registros, {len(self.isbn_index)} ISBN indexados, {len(self.issn_index)} ISSN indexados")
    def _read_headers(self, ws):
        return [clean_cell(cell.value) or "" for cell in ws[1]]
    def _validate_headers(self, headers):
        normalized = {normalize_header(h) for h in headers}
        missing = [c for c in self.REQUIRED_COLUMNS if c not in normalized]
        if missing:
            raise ValueError("El Excel de Alma no contiene las columnas esperadas. Faltan: " + ", ".join(missing))
    def find_by_isbn(self, isbn):
        return self.isbn_index.get(normalize_identifier(isbn), []) if normalize_identifier(isbn) else []
    def find_by_issn(self, issn):
        return self.issn_index.get(normalize_identifier(issn), []) if normalize_identifier(issn) else []
    def find_by_title(self, title):
        if not title:
            return {"status": "NO_MATCH", "candidates": [], "scores": []}
        scored = []
        for rec in self.records:
            if rec.normalized_title_words:
                score = title_similarity_score(title, rec)
                if score >= TITLE_MATCH_MIN_SCORE:
                    scored.append((rec, score))
        scored.sort(key=lambda x: x[1], reverse=True)
        if not scored:
            return {"status": "NO_MATCH", "candidates": [], "scores": []}
        if len(scored) == 1 or scored[0][1] - scored[1][1] >= TITLE_MATCH_MIN_MARGIN:
            return {"status": "MATCH", "candidates": [scored[0][0]], "scores": [scored[0][1]]}
        best = scored[0][1]
        amb = [(r, s) for r, s in scored if best - s < TITLE_MATCH_MIN_MARGIN]
        return {"status": "AMBIGUOUS", "candidates": [r for r, s in amb], "scores": [s for r, s in amb]}

class FacturaeParser:
    def __init__(self, invoice_path):
        self.invoice_path = invoice_path
        self.tree = None
        self.root = None
    def load(self):
        try:
            self.tree = ET.parse(self.invoice_path)
            self.root = self.tree.getroot()
            logging.info(f"Archivo de factura cargado correctamente: {self.invoice_path}")
            return
        except DefusedXmlException as e:
            logging.exception("XML bloqueado por seguridad")
            raise ValueError("El archivo XML/XSIG contiene estructuras XML no permitidas por seguridad y no se procesara. " + str(e))
        except Exception as first_error:
            logging.warning(f"No se pudo parsear directamente. Se intentara lectura alternativa. Error inicial: {first_error}")
        encodings = ["utf-8-sig", "utf-8", "iso-8859-1", "latin-1"]
        last_error = None
        for enc in encodings:
            try:
                with open(self.invoice_path, "r", encoding=enc) as f:
                    content = f.read().strip()
                xml_start = content.find("<")
                if xml_start > 0:
                    content = content[xml_start:]
                self.root = ET.fromstring(content)
                self.tree = ET.ElementTree(self.root)
                logging.info(f"Archivo de factura cargado mediante lectura alternativa con codificacion {enc}: {self.invoice_path}")
                return
            except DefusedXmlException as e:
                logging.exception("XML bloqueado por seguridad")
                raise ValueError("El archivo XML/XSIG contiene estructuras XML no permitidas por seguridad y no se procesara. " + str(e))
            except Exception as e:
                last_error = e
        logging.exception("Error cargando archivo de factura")
        raise ValueError(f"No se pudo leer o parsear el archivo de factura. Debe ser XML valido. Ultimo error: {last_error}")
    def find_text(self, element, tag, default=""):
        if element is None:
            return default
        found = element.find(f".//{{*}}{tag}")
        return found.text.strip() if found is not None and found.text else default
    def find_direct_text(self, element, tag, default=""):
        if element is None:
            return default
        found = element.find(f"./{{*}}{tag}")
        return found.text.strip() if found is not None and found.text else default
    def find_path_text(self, element, path, default=""):
        if element is None:
            return default
        xpath = "./" + "/".join(f"{{*}}{part}" for part in path.split("/"))
        found = element.find(xpath)
        return found.text.strip() if found is not None and found.text else default
    def get_seller_party_name(self, seller_party):
        if seller_party is None:
            return None
        legal = seller_party.find(".//{*}LegalEntity")
        individual = seller_party.find(".//{*}Individual")
        for candidate in [self.find_direct_text(legal, "CorporateName"), self.find_direct_text(individual, "CorporateName"), self.find_direct_text(legal, "TradeName"), self.find_direct_text(individual, "TradeName")]:
            if candidate:
                return candidate
        parts = [self.find_direct_text(individual, "Name"), self.find_direct_text(individual, "FirstSurname"), self.find_direct_text(individual, "SecondSurname")]
        full = " ".join(p.strip() for p in parts if p and p.strip())
        return full if full else None
    def parse(self):
        if self.root is None:
            self.load()
        invoices_xml = self.root.findall(".//{*}Invoice")
        logging.info(f"Numero de nodos Invoice detectados: {len(invoices_xml)}")
        if not invoices_xml:
            raise ValueError("No se encontraron elementos <Invoice> en el archivo.")
        return [self._parse_invoice(inv) for inv in invoices_xml]
    def _parse_invoice(self, inv):
        seller_party = self.root.find(".//{*}SellerParty")
        seller_tax_id = self.find_text(seller_party, "TaxIdentificationNumber", "")
        seller_name = self.get_seller_party_name(seller_party)
        vendor_account = get_vendor_account(seller_tax_id, seller_name)
        invoice_data = InvoiceData(
            invoice_number=self.find_text(inv, "InvoiceNumber"),
            date=parse_date_for_excel(self.find_text(inv, "IssueDate")),
            total_amount=safe_decimal(self.find_text(inv, "InvoiceTotal")),
            currency=self.find_text(inv, "InvoiceCurrencyCode", "EUR"),
            vendor_account=vendor_account,
            note=self.find_text(inv, "InvoiceDescription") or None,
            seller_name=seller_name,
            seller_tax_id=seller_tax_id,
            source_file=self.invoice_path,
            vat_amount=decimal_or_none(self.find_path_text(inv, "TaxesOutputs/Tax/TaxAmount/TotalAmount")),
            inclusive=False,
            lines=[]
        )
        lines_xml = inv.findall(".//{*}InvoiceLine")
        logging.info(f"Factura {invoice_data.invoice_number}: {len(lines_xml)} lineas")
        for idx, line_xml in enumerate(lines_xml, start=1):
            invoice_data.lines.append(self._parse_line(line_xml, idx))
        return invoice_data
    def _parse_line(self, line, idx):
        seq_str = self.find_text(line, "SequenceNumber", str(idx))
        try:
            seq_num = int(seq_str)
        except ValueError:
            seq_num = idx
        raw_title = self.find_text(line, "ItemDescription") or None
        desc_isbn, desc_issn, cleaned_title = extract_identifiers_from_text(raw_title)
        title = cleaned_title or raw_title
        quantity = safe_decimal(self.find_text(line, "Quantity", "1"), Decimal("1.00"))
        gross_amount_text = self.find_path_text(line, "GrossAmount")
        unit_price_text = self.find_text(line, "UnitPriceWithoutTax")
        price = safe_decimal(gross_amount_text, Decimal("0.00")) if gross_amount_text else safe_decimal(unit_price_text, Decimal("0.00"))
        issn, isbn, vendor_ref = classify_article_code(self.find_text(line, "ArticleCode") or None)
        if not isbn and desc_isbn:
            isbn = desc_isbn
        if not issn and desc_issn:
            issn = desc_issn
        return InvoiceLineData(
            line_number=seq_num,
            po_line=None,
            vendor_ref=vendor_ref,
            issn=issn,
            isbn=isbn,
            title=title,
            price=price,
            quantity=quantity,
            vat_amount=decimal_or_none(self.find_path_text(line, "TaxesOutputs/Tax/TaxAmount/TotalAmount")),
            vat_percentage=decimal_or_none(self.find_path_text(line, "TaxesOutputs/Tax/TaxRate"))
        )

class AlmaLineEnricher:
    def __init__(self, lookup):
        self.lookup = lookup
        self.validation_issues = []
    def enrich(self, invoices):
        for invoice in invoices:
            for line in invoice.lines:
                self._enrich_line(invoice, line)
        return self.validation_issues
    def _clear_alma_fields(self, line):
        line.po_line = None
        line.reporting_code = None
        line.secondary_reporting_code = None
        line.tertiary_reporting_code = None
        line.fourth_reporting_code = None
        line.fifth_reporting_code = None
        line.funds = []
    def _enrich_line(self, invoice, line):
        if line.isbn:
            c = self.lookup.find_by_isbn(line.isbn)
            if c:
                self._resolve_candidates(invoice, line, c, "ISBN")
                return
        if line.issn:
            c = self.lookup.find_by_issn(line.issn)
            if c:
                self._resolve_candidates(invoice, line, c, "ISSN")
                return
        if line.title:
            result = self.lookup.find_by_title(line.title)
            if result["status"] == "MATCH":
                self._resolve_candidates(invoice, line, result["candidates"], "TITLE")
                return
            if result["status"] == "AMBIGUOUS":
                candidate_po_lines = [r.po_line for r in result["candidates"]]
                score_text = "; ".join(f"{r.po_line}: {s:.3f}" for r, s in zip(result["candidates"], result["scores"]))
                self._clear_alma_fields(line)
                line.match_status = "AMBIGUOUS_TITLE"
                line.match_message = "Varias coincidencias posibles por titulo"
                line.match_candidates = candidate_po_lines
                self.validation_issues.append(ValidationIssue("TITULO AMBIGUO", invoice.invoice_number, line.line_number, line.isbn, line.issn, line.title, candidate_po_lines, "Varias PO Lines candidatas por titulo. Puntuaciones: " + score_text))
                return
        self._clear_alma_fields(line)
        line.match_status = "NO_MATCH"
        line.match_message = "Sin coincidencia en Alma por ISBN, ISSN ni titulo"
        self.validation_issues.append(ValidationIssue("SIN COINCIDENCIA", invoice.invoice_number, line.line_number, line.isbn, line.issn, line.title, [], "No se encontro ninguna PO Line por ISBN, ISSN ni titulo."))
    def _resolve_candidates(self, invoice, line, candidates, match_type):
        unique = {c.po_line: c for c in candidates}
        unique_candidates = list(unique.values())
        if len(unique_candidates) == 1:
            r = unique_candidates[0]
            line.po_line = r.po_line
            line.reporting_code = r.reporting_code
            line.secondary_reporting_code = r.secondary_reporting_code
            line.tertiary_reporting_code = r.tertiary_reporting_code
            line.fourth_reporting_code = r.fourth_reporting_code
            line.fifth_reporting_code = r.fifth_reporting_code
            line.alma_quantity_for_pricing = r.quantity_for_pricing
            line.funds = list(r.funds or [])
            if r.title:
                line.title = r.title
            if r.start_subs_date:
                line.start_date = r.start_subs_date
            line.match_status = f"MATCH_{match_type}"
            line.match_message = f"Coincidencia unica por {match_type}"
            line.match_candidates = [r.po_line]

            if (
                r.quantity_for_pricing is not None
                and line.quantity is not None
                and line.quantity != r.quantity_for_pricing
            ):
                self.validation_issues.append(
                    ValidationIssue(
                        issue_type="REVISAR CANTIDAD",
                        invoice_number=invoice.invoice_number,
                        invoice_line_number=line.line_number,
                        isbn=line.isbn,
                        issn=line.issn,
                        title=line.title,
                        candidates=[r.po_line],
                        message=(
                            f"La factura indica Quantity={line.quantity}, mientras que "
                            f"Alma indica Quantity for Pricing={r.quantity_for_pricing}. "
                            "Puede tratarse de una facturacion parcial o de una entrega final "
                            "que complete una linea facturada previamente. "
                            "Revisar manualmente antes de cargar en Alma."
                        ),
                        po_line=r.po_line,
                        invoice_quantity=line.quantity,
                        alma_quantity_for_pricing=r.quantity_for_pricing,
                    )
                )
            return
        candidate_po_lines = sorted(r.po_line for r in unique_candidates)
        self._clear_alma_fields(line)
        line.match_status = f"DUPLICATE_{match_type}"
        line.match_message = f"Duplicado por {match_type}. Revision manual necesaria."
        line.match_candidates = candidate_po_lines
        self.validation_issues.append(ValidationIssue(f"DUPLICADO {match_type}", invoice.invoice_number, line.line_number, line.isbn, line.issn, line.title, candidate_po_lines, "El identificador tiene varias PO Lines candidatas."))

class AlmaExcelExporter:
    DATE_FORMAT = "mm/dd/yyyy"
    def __init__(self, output_path, template_path=None):
        self.output_path = output_path
        self.template_path = template_path
        self.workbook = None
        self.worksheet = None
        self.hinv_headers = []
        self.hil_headers = []
        self.hinv_style = None
        self.inv_style = None
        self.hil_style = None
        self.il_style = None
        self.column_widths = {}
    def create_or_load_workbook(self):
        if self.template_path and os.path.exists(self.template_path):
            self._load_from_template()
        else:
            self._create_without_template()
    def _load_from_template(self):
        wb = load_workbook(self.template_path)
        ws = wb.active
        hinv_row = self._find_marker_row(ws, "HINV")
        hil_row = self._find_marker_row(ws, "HIL")
        if not hinv_row or not hil_row:
            raise ValueError("La plantilla seleccionada no contiene filas HINV y HIL reconocibles.")
        inv_row = hinv_row + 1
        il_row = hil_row + 1
        self.hinv_headers = self._read_row_values(ws, hinv_row)
        self.hil_headers = self._read_row_values(ws, hil_row)
        max_cols = max(len(self.hinv_headers), len(self.hil_headers), 25)
        self.hinv_style = self._capture_row_style(ws, hinv_row, max_cols)
        self.inv_style = self._capture_row_style(ws, inv_row, max_cols)
        self.hil_style = self._capture_row_style(ws, hil_row, max_cols)
        self.il_style = self._capture_row_style(ws, il_row, max_cols)
        self.column_widths = {cl: ws.column_dimensions[cl].width for cl in ws.column_dimensions}
        self.workbook = wb
        self.worksheet = ws
        self._delete_other_sheets()
        self._clear_sheet_fully(ws)
        for cl, width in self.column_widths.items():
            if width:
                self.worksheet.column_dimensions[cl].width = width
    def _create_without_template(self):
        self.workbook = Workbook()
        self.worksheet = self.workbook.active
        self.worksheet.title = "Sheet1"
        self.hinv_headers = DEFAULT_HINV_HEADERS
        self.hil_headers = DEFAULT_HIL_HEADERS
    def _find_marker_row(self, ws, marker):
        for r in range(1, ws.max_row + 1):
            if clean_cell(ws.cell(r, 1).value) == marker:
                return r
        return None
    def _read_row_values(self, ws, row_idx):
        last_col = 0
        for c in range(1, ws.max_column + 1):
            if clean_cell(ws.cell(row_idx, c).value) is not None:
                last_col = c
        return [clean_cell(ws.cell(row_idx, c).value) or "" for c in range(1, last_col + 1)] if last_col else []
    def _capture_row_style(self, ws, row_idx, max_cols):
        styles = []
        for c in range(1, max_cols + 1):
            cell = ws.cell(row_idx, c)
            styles.append({"style": copy.copy(cell._style) if cell.has_style else None, "number_format": cell.number_format, "font": copy.copy(cell.font), "fill": copy.copy(cell.fill), "border": copy.copy(cell.border), "alignment": copy.copy(cell.alignment), "protection": copy.copy(cell.protection)})
        return {"height": ws.row_dimensions[row_idx].height, "cells": styles}
    def _apply_row_style(self, row_idx, row_style, max_cols):
        if not row_style:
            return
        self.worksheet.row_dimensions[row_idx].height = row_style["height"]
        for c in range(1, min(max_cols, len(row_style["cells"])) + 1):
            info = row_style["cells"][c - 1]
            cell = self.worksheet.cell(row_idx, c)
            if info["style"] is not None:
                cell._style = copy.copy(info["style"])
            cell.number_format = info["number_format"]
            cell.font = copy.copy(info["font"])
            cell.fill = copy.copy(info["fill"])
            cell.border = copy.copy(info["border"])
            cell.alignment = copy.copy(info["alignment"])
            cell.protection = copy.copy(info["protection"])
    def _delete_other_sheets(self):
        active = self.worksheet
        for ws in list(self.workbook.worksheets):
            if ws != active:
                self.workbook.remove(ws)
    def _clear_sheet_fully(self, ws):
        if ws.max_row > 0:
            ws.delete_rows(1, ws.max_row)
    def _write_row(self, row_idx, values, row_style):
        self._apply_row_style(row_idx, row_style, len(values))
        for c, value in enumerate(values, start=1):
            cell = self.worksheet.cell(row_idx, c)
            value = sanitize_excel_value(value)
            cell.value = value
            if isinstance(value, datetime):
                cell.number_format = self.DATE_FORMAT
    def _find_header_col(self, headers, target_header):
        target = normalize_header(target_header)
        for idx, header in enumerate(headers, start=1):
            if normalize_header(header) == target:
                return idx
        return None
    def _apply_line_type_dropdown(self, start_row, end_row):
        if end_row < start_row:
            return
        col = self._find_header_col(self.hil_headers, "Line type")
        if not col:
            return
        validation = DataValidation(type="list", formula1='"' + ",".join(LINE_TYPE_OPTIONS) + '"', allow_blank=False)
        validation.error = "Selecciona un tipo de linea valido para Alma."
        validation.errorTitle = "Tipo de linea no valido"
        validation.prompt = "Selecciona un tipo de linea."
        validation.promptTitle = "Line type"
        self.worksheet.add_data_validation(validation)
        validation.add(f"{get_column_letter(col)}{start_row}:{get_column_letter(col)}{end_row}")
    def write_invoices(self, invoices):
        if self.workbook is None or self.worksheet is None:
            self.create_or_load_workbook()
        row = 1
        for invoice in invoices:
            self._write_row(row, self.hinv_headers, self.hinv_style); row += 1
            self._write_row(row, self._build_invoice_row_from_headers(invoice, self.hinv_headers), self.inv_style); row += 1
            self._write_row(row, self.hil_headers, self.hil_style); row += 1
            first_il_row = row
            for line in invoice.lines:
                self._write_row(row, self._build_line_row_from_headers(line, self.hil_headers), self.il_style)
                row += 1
            self._apply_line_type_dropdown(first_il_row, row - 1)
    def _build_invoice_row_from_headers(self, invoice, headers):
        values = {"HINV": "INV", "Invoice Number": invoice.invoice_number, "Date": invoice.date, "Total amount": float(invoice.total_amount), "Currency": invoice.currency, "Vendor Account": invoice.vendor_account, "Payment method": None, "Note": invoice.note, "Use Pro Rata": None, "Shipment amount": None, "Overhead amount": None, "Insurance amount": None, "Discount amount": None, "VAT Percentage": None, "VAT Amount": decimal_to_float_or_none(invoice.vat_amount), "Inclusive": invoice.inclusive, "Expended From Fund": None, "PrePaid": None, "Payment status": None, "Voucher number": None, "Voucher date": None, "Voucher Amount": None, "Invoice Reference Number": None, "VAT In Invoice Line Level": "YES", "Report TAX": None, "VAT code": None}
        return [values.get(normalize_header(h), None) for h in headers]
    def _build_line_row_from_headers(self, line, headers):
        values = {"HIL": "IL", "Line type": "REGULAR", "Line Number": line.line_number, "PO Line": line.po_line, "Vendor Ref Number": None, "ISSN": line.issn, "ISBN": line.isbn, "Title": line.title, "Price": float(line.price), "Price Note": None, "Quantity": float(line.quantity), "Reporting Code": line.reporting_code, "Secondary Reporting Code": line.secondary_reporting_code, "Tertiary Reporting Code": line.tertiary_reporting_code, "Fourth Reporting Code": line.fourth_reporting_code, "Fifth Reporting Code": line.fifth_reporting_code, "VAT code": None, "Start subs date": None, "End subs date": None, "Additional Info": None, "Note": None, "VAT Amount": decimal_to_float_or_none(line.vat_amount), "VAT Percentage": decimal_to_float_or_none(line.vat_percentage)}
        row = []
        fund_index = 0
        for h in headers:
            nh = normalize_header(h)
            if nh == "Fund and percent":
                row.append(line.funds[fund_index] if fund_index < len(line.funds) else None)
                fund_index += 1
            else:
                row.append(values.get(nh, None))
        return row
    def save(self):
        try:
            self.workbook.save(self.output_path)
            logging.info(f"Excel Alma guardado correctamente: {self.output_path}")
        except PermissionError:
            logging.exception("No se puede guardar el Excel porque esta abierto o bloqueado")
            raise ValueError(f"No se pudo guardar el archivo '{self.output_path}'. Puede que este abierto en Excel.")
        except Exception as e:
            logging.exception("Error guardando Excel")
            raise ValueError(f"Error al guardar el Excel: {e}")

class ValidationReportExporter:
    def __init__(self, output_path):
        self.output_path = output_path
    def save(self, issues):
        if not issues:
            return None
        wb = Workbook()
        ws = wb.active
        ws.title = "Validacion"
        ws.append([
            "Tipo aviso", "Factura", "Linea factura", "PO Line",
            "ISBN XSIG/XML", "ISSN XSIG/XML", "Titulo",
            "Cantidad factura", "Quantity for Pricing Alma",
            "PO Lines candidatas", "Mensaje"
        ])
        for issue in issues:
            ws.append([sanitize_excel_value(v) for v in [
                issue.issue_type,
                issue.invoice_number,
                issue.invoice_line_number,
                issue.po_line,
                issue.isbn,
                issue.issn,
                issue.title,
                decimal_to_float_or_none(issue.invoice_quantity),
                decimal_to_float_or_none(issue.alma_quantity_for_pricing),
                "; ".join(issue.candidates),
                issue.message,
            ]])
        wb.save(self.output_path)
        logging.info(f"Informe de validacion guardado: {self.output_path}")
        return self.output_path

def parse_invoice_paths(invoice_paths):
    if isinstance(invoice_paths, str):
        invoice_paths = [invoice_paths]
    invoices = []
    for p in invoice_paths:
        invoices.extend(FacturaeParser(p).parse())
    return invoices

def validate_single_vendor_batch(invoices):
    keys = {}
    for inv in invoices:
        key = get_vendor_batch_key(inv)
        label = inv.vendor_account or inv.seller_name or "Proveedor no identificado"
        keys.setdefault(key, set()).add(label)
    meaningful_keys = {k for k in keys if k}
    if len(meaningful_keys) <= 1:
        return
    details = []
    for labels in keys.values():
        for label in sorted(labels):
            details.append(f"- {label}")
    raise ValueError("El lote contiene facturas de varios proveedores. Selecciona unicamente facturas del mismo proveedor.\n\nProveedores detectados:\n" + "\n".join(details))

def convert_facturae_to_alma_excel(invoice_paths, alma_pol_excel_path, excel_output_path, template_path=None):
    invoices = parse_invoice_paths(invoice_paths)
    if not invoices:
        raise ValueError("No se encontro ninguna factura para procesar.")
    validate_single_vendor_batch(invoices)
    lookup = AlmaPOLineLookup(alma_pol_excel_path)
    lookup.load()
    enricher = AlmaLineEnricher(lookup)
    validation_issues = enricher.enrich(invoices)
    exporter = AlmaExcelExporter(output_path=excel_output_path, template_path=template_path)
    exporter.create_or_load_workbook()
    exporter.write_invoices(invoices)
    exporter.save()
    validation_path = None
    if validation_issues:
        validation_path = build_validation_output_path(excel_output_path)
        ValidationReportExporter(validation_path).save(validation_issues)
    summary = build_conversion_summary(invoices, validation_issues)
    return len(invoices), len(validation_issues), validation_path, summary

class FacturaeApp:
    DEFAULT_INSTITUTION = "Universidad de Valladolid. Biblioteca"

    def __init__(self, root):
        self.root = root
        self.config_path = self._get_config_path()
        self.preferences = self._load_preferences()
        self.conversion_mode = tk.StringVar(value="single")
        self.invoice_file = tk.StringVar()
        self.alma_pol_file = tk.StringVar()
        self.template_file = tk.StringVar()
        self.output_file = tk.StringVar()
        self.status_text = tk.StringVar(value="Estado: pendiente de seleccionar archivos.")
        self.institution_var = tk.StringVar(value=self.preferences.get("institution", self.DEFAULT_INSTITUTION))
        self.custom_invoice_dir = tk.StringVar(value=self.preferences.get("custom_invoice_dir", ""))
        self.custom_alma_dir = tk.StringVar(value=self.preferences.get("custom_alma_dir", ""))
        self.custom_template_dir = tk.StringVar(value=self.preferences.get("custom_template_dir", ""))
        self.custom_output_dir = tk.StringVar(value=self.preferences.get("custom_output_dir", ""))
        self.logo_path_var = tk.StringVar(value=self.preferences.get("logo_path", ""))
        self.logo_photo = None
        self.invoice_files = []
        self.invoice_metadata = {}
        self.worker_queue = queue.Queue()
        self.worker_thread = None
        remembered_template = self.preferences.get("last_template_path", "")
        if remembered_template and os.path.isfile(remembered_template):
            self.template_file.set(remembered_template)
        self.root.geometry("1180x940")
        self.root.minsize(1000, 820)
        self.root.resizable(True, True)
        self.root.configure(bg=UI_BG)
        self._setup_style()
        self._build_ui()
        self._update_title()
        self._load_logo_from_preferences(show_warning=False)
        self.root.protocol("WM_DELETE_WINDOW", self._on_close)

    def _application_dir(self):
        if getattr(sys, "frozen", False):
            return os.path.dirname(sys.executable)
        return os.path.dirname(os.path.abspath(__file__))

    def _get_config_path(self):
        base = os.environ.get("APPDATA") or os.path.expanduser("~")
        folder = os.path.join(base, APP_NAME)
        try:
            os.makedirs(folder, exist_ok=True)
            return os.path.join(folder, "config.json")
        except OSError:
            return os.path.join(self._application_dir(), "facturae_to_alma_config.json")

    def _default_preferences(self):
        return {
            "institution": self.DEFAULT_INSTITUTION,
            "custom_invoice_dir": "",
            "custom_alma_dir": "",
            "custom_template_dir": "",
            "custom_output_dir": "",
            "logo_path": "",
            "last_invoice_dir": "",
            "last_alma_dir": "",
            "last_template_dir": "",
            "last_output_dir": "",
            "last_template_path": "",
        }

    def _load_preferences(self):
        data = self._default_preferences()
        try:
            if os.path.isfile(self.config_path):
                with open(self.config_path, "r", encoding="utf-8") as f:
                    loaded = json.load(f)
                if isinstance(loaded, dict):
                    data.update(loaded)
        except (OSError, json.JSONDecodeError) as e:
            logging.warning(f"No se pudo leer la configuracion: {e}")
        return data

    def _save_preferences_file(self):
        try:
            os.makedirs(os.path.dirname(self.config_path), exist_ok=True)
            with open(self.config_path, "w", encoding="utf-8") as f:
                json.dump(self.preferences, f, ensure_ascii=False, indent=2)
        except OSError as e:
            logging.exception("No se pudo guardar la configuracion")
            raise ValueError(f"No se pudo guardar la configuracion: {e}")

    def _valid_dir(self, value):
        return value if value and os.path.isdir(value) else ""

    def _initial_dir(self, custom_key, last_key, fallback=None):
        custom = self._valid_dir(self.preferences.get(custom_key, ""))
        if custom:
            return custom
        last = self._valid_dir(self.preferences.get(last_key, ""))
        if last:
            return last
        if fallback and os.path.isdir(fallback):
            return fallback
        return self._application_dir()

    def _remember_dir(self, key, path):
        folder = path if os.path.isdir(path) else os.path.dirname(path)
        if folder and os.path.isdir(folder):
            self.preferences[key] = folder
            try:
                self._save_preferences_file()
            except ValueError:
                logging.warning("No se pudo persistir el ultimo directorio")

    def _setup_style(self):
        style = ttk.Style()
        try:
            style.theme_use("clam")
        except tk.TclError as exc:
            logging.warning("No se pudo aplicar el tema grafico 'clam': %s", exc)
        style.configure("App.TFrame", background=UI_BG)
        style.configure("App.TLabel", background=UI_BG)
        style.configure("App.TLabelframe", background=UI_BG)
        style.configure("App.TLabelframe.Label", background=UI_BG)
        style.configure("App.TRadiobutton", background=UI_BG)
        style.configure("Primary.TButton", font=("Segoe UI", 11, "bold"), padding=(20, 12))
        style.configure("HelpTitle.TLabel", background=UI_BG, font=("Segoe UI", 12, "bold"))
        style.configure("Treeview", rowheight=25)

    def _update_title(self, *_):
        institution = self.institution_var.get().strip() or self.DEFAULT_INSTITUTION
        title = f"{institution} - {APP_NAME} {APP_VERSION}"
        self.root.title(title)
        if hasattr(self, "title_label"):
            self.title_label.config(text=title)

    def _build_ui(self):
        shell = ttk.Frame(self.root, padding=16, style="App.TFrame")
        shell.pack(fill=tk.BOTH, expand=True)
        self.header_frame = ttk.Frame(shell, style="App.TFrame")
        self.header_frame.pack(fill=tk.X, pady=(0, 12))
        self.logo_label = ttk.Label(self.header_frame, style="App.TLabel")
        self.logo_label.pack(anchor="w", pady=(0, 8))
        self.title_label = ttk.Label(
            self.header_frame,
            text="",
            font=("Segoe UI", 15, "bold"),
            style="App.TLabel",
        )
        self.title_label.pack(fill=tk.X, anchor="w")
        self.notebook = ttk.Notebook(shell)
        self.notebook.pack(fill=tk.BOTH, expand=True)
        self.convert_tab = ttk.Frame(self.notebook, padding=14, style="App.TFrame")
        self.customize_tab = ttk.Frame(self.notebook, padding=18, style="App.TFrame")
        self.help_tab = ttk.Frame(self.notebook, padding=18, style="App.TFrame")
        self.notebook.add(self.convert_tab, text="Convertir")
        self.notebook.add(self.customize_tab, text="Personalizar")
        self.notebook.add(self.help_tab, text="Ayuda")
        self._build_convert_tab()
        self._build_customize_tab()
        self._build_help_tab()

    def _build_convert_tab(self):
        tab = self.convert_tab
        tab.columnconfigure(0, weight=1)
        mode = ttk.LabelFrame(tab, text="1. Modo de conversion", padding=10, style="App.TLabelframe")
        mode.grid(row=0, column=0, sticky="ew")
        ttk.Radiobutton(mode, text="Factura individual", variable=self.conversion_mode, value="single", command=self.on_mode_change, style="App.TRadiobutton").grid(row=0, column=0, padx=(0, 24), sticky="w")
        ttk.Radiobutton(mode, text="Lote de facturas del mismo proveedor y biblioteca", variable=self.conversion_mode, value="batch", command=self.on_mode_change, style="App.TRadiobutton").grid(row=0, column=1, sticky="w")

        input_frame = ttk.LabelFrame(tab, text="2. Archivos de entrada", padding=10, style="App.TLabelframe")
        input_frame.grid(row=1, column=0, sticky="ew", pady=(10, 0))
        input_frame.columnconfigure(0, weight=1)
        self.invoice_label = ttk.Label(input_frame, text="Factura electronica XSIG / XML / TXT *", style="App.TLabel")
        self.invoice_label.grid(row=0, column=0, sticky="w")
        ttk.Entry(input_frame, textvariable=self.invoice_file, state="readonly").grid(row=1, column=0, sticky="ew", padx=(0, 10), pady=(2, 8))
        self.invoice_button = ttk.Button(input_frame, text="Seleccionar factura...", command=self.select_invoice_or_batch)
        self.invoice_button.grid(row=1, column=1, sticky="ew", pady=(2, 8))
        ttk.Label(input_frame, text="Fichero Excel de Alma Analytics con PO Lines *", style="App.TLabel").grid(row=2, column=0, sticky="w")
        ttk.Entry(input_frame, textvariable=self.alma_pol_file).grid(row=3, column=0, sticky="ew", padx=(0, 10), pady=(2, 8))
        ttk.Button(input_frame, text="Seleccionar fichero...", command=self.select_alma_pol).grid(row=3, column=1, sticky="ew", pady=(2, 8))
        ttk.Label(input_frame, text="Plantilla Alma .xlsx *", style="App.TLabel").grid(row=4, column=0, sticky="w")
        ttk.Entry(input_frame, textvariable=self.template_file).grid(row=5, column=0, sticky="ew", padx=(0, 10), pady=(2, 0))
        ttk.Button(input_frame, text="Seleccionar plantilla...", command=self.select_template).grid(row=5, column=1, sticky="ew", pady=(2, 0))

        self.batch_frame = ttk.LabelFrame(tab, text="Facturas del lote", padding=8, style="App.TLabelframe")
        self.batch_frame.grid(row=2, column=0, sticky="nsew", pady=(10, 0))
        self.batch_frame.columnconfigure(0, weight=1)
        self.batch_frame.rowconfigure(0, weight=1)
        cols = ("invoice", "vendor", "file")
        self.batch_tree = ttk.Treeview(self.batch_frame, columns=cols, show="headings", height=10, selectmode="extended")
        self.batch_tree.heading("invoice", text="Numero de factura")
        self.batch_tree.heading("vendor", text="Proveedor")
        self.batch_tree.heading("file", text="Fichero")
        self.batch_tree.column("invoice", width=150, anchor="w")
        self.batch_tree.column("vendor", width=310, anchor="w")
        self.batch_tree.column("file", width=520, anchor="w")
        self.batch_tree.grid(row=0, column=0, sticky="nsew")
        scroll = ttk.Scrollbar(self.batch_frame, orient="vertical", command=self.batch_tree.yview)
        scroll.grid(row=0, column=1, sticky="ns")
        self.batch_tree.configure(yscrollcommand=scroll.set)
        batch_buttons = ttk.Frame(self.batch_frame, style="App.TFrame")
        batch_buttons.grid(row=1, column=0, columnspan=2, sticky="w", pady=(7, 0))
        ttk.Button(batch_buttons, text="Anadir facturas", command=self.add_batch_files).pack(side=tk.LEFT, padx=(0, 6))
        ttk.Button(batch_buttons, text="Quitar seleccionadas", command=self.remove_selected_batch_files).pack(side=tk.LEFT, padx=6)
        ttk.Button(batch_buttons, text="Ordenar por numero", command=self.sort_batch_files).pack(side=tk.LEFT, padx=6)
        ttk.Button(batch_buttons, text="Vaciar lote", command=self.clear_batch_files).pack(side=tk.LEFT, padx=6)

        output = ttk.LabelFrame(tab, text="3. Archivo de salida", padding=10, style="App.TLabelframe")
        output.grid(row=3, column=0, sticky="ew", pady=(10, 0))
        output.columnconfigure(0, weight=1)
        ttk.Entry(output, textvariable=self.output_file).grid(row=0, column=0, sticky="ew", padx=(0, 10))
        ttk.Button(output, text="Elegir destino...", command=self.select_output).grid(row=0, column=1)

        actions = ttk.LabelFrame(tab, text="4. Acciones", padding=10, style="App.TLabelframe")
        actions.grid(row=4, column=0, sticky="ew", pady=(10, 0))
        actions.columnconfigure(0, weight=2)
        for c in (1, 2, 3): actions.columnconfigure(c, weight=1)
        self.convert_button = ttk.Button(actions, text="Convertir a formato Alma", command=self.run_conversion, style="Primary.TButton")
        self.convert_button.grid(row=0, column=0, sticky="ew", padx=(0, 8))
        ttk.Button(actions, text="Limpiar datos", command=self.clear_fields).grid(row=0, column=1, sticky="ew", padx=4)
        ttk.Button(actions, text="Abrir carpeta de salida", command=self.open_output_folder).grid(row=0, column=2, sticky="ew", padx=4)
        ttk.Button(actions, text="Salir", command=self._on_close).grid(row=0, column=3, sticky="ew", padx=(4, 0))
        ttk.Label(tab, textvariable=self.status_text, style="App.TLabel", font=("Segoe UI", 9, "italic")).grid(row=5, column=0, sticky="w", pady=(10, 0))
        tab.rowconfigure(2, weight=1)
        self.on_mode_change(initial=True)

    def _directory_row(self, parent, row, label, variable, command):
        ttk.Label(parent, text=label, style="App.TLabel").grid(row=row, column=0, sticky="w", pady=6)
        ttk.Entry(parent, textvariable=variable).grid(row=row, column=1, sticky="ew", padx=10, pady=6)
        ttk.Button(parent, text="Seleccionar carpeta...", command=command).grid(row=row, column=2, pady=6)

    def _build_customize_tab(self):
        tab = self.customize_tab
        tab.columnconfigure(0, weight=1)
        frame = ttk.LabelFrame(tab, text="Preferencias", padding=16, style="App.TLabelframe")
        frame.grid(row=0, column=0, sticky="nsew")
        frame.columnconfigure(1, weight=1)
        ttk.Label(frame, text="Institucion", style="App.TLabel").grid(row=0, column=0, sticky="w", pady=6)
        ttk.Entry(frame, textvariable=self.institution_var).grid(row=0, column=1, columnspan=2, sticky="ew", padx=(10, 0), pady=6)
        self._directory_row(frame, 1, "Directorio de facturas FACe", self.custom_invoice_dir, lambda: self._choose_preference_dir(self.custom_invoice_dir))
        self._directory_row(frame, 2, "Directorio del fichero Excel de Alma", self.custom_alma_dir, lambda: self._choose_preference_dir(self.custom_alma_dir))
        self._directory_row(frame, 3, "Directorio de plantilla Alma", self.custom_template_dir, lambda: self._choose_preference_dir(self.custom_template_dir))
        self._directory_row(frame, 4, "Directorio de Excel finales", self.custom_output_dir, lambda: self._choose_preference_dir(self.custom_output_dir))
        ttk.Label(frame, text="Logo institucional", style="App.TLabel").grid(row=5, column=0, sticky="w", pady=6)
        ttk.Entry(frame, textvariable=self.logo_path_var).grid(row=5, column=1, sticky="ew", padx=10, pady=6)
        logo_buttons = ttk.Frame(frame, style="App.TFrame")
        logo_buttons.grid(row=5, column=2, sticky="e", pady=6)
        ttk.Button(logo_buttons, text="Seleccionar logo...", command=self.select_logo).pack(side=tk.LEFT, padx=(0, 6))
        ttk.Button(logo_buttons, text="Eliminar logo", command=self.remove_logo).pack(side=tk.LEFT)
        self.logo_status_label = ttk.Label(frame, text="PNG, JPG o JPEG. Altura maxima: 100 px.", style="App.TLabel")
        self.logo_status_label.grid(row=6, column=1, columnspan=2, sticky="w", padx=10)
        note = ("Todos los campos son opcionales. Si una carpeta personalizada queda vacia, se usa la ultima carpeta valida; "
                "en la primera ejecucion se usa la carpeta de la aplicacion. La salida se guarda junto a la factura, salvo que se defina un directorio de salida.")
        ttk.Label(frame, text=note, wraplength=880, justify="left", style="App.TLabel").grid(row=7, column=0, columnspan=3, sticky="w", pady=(14, 8))
        buttons = ttk.Frame(frame, style="App.TFrame")
        buttons.grid(row=8, column=0, columnspan=3, sticky="w", pady=(10, 0))
        ttk.Button(buttons, text="Guardar preferencias", command=self.save_preferences).pack(side=tk.LEFT, padx=(0, 8))
        ttk.Button(buttons, text="Restablecer valores", command=self.reset_preferences).pack(side=tk.LEFT)

    def _build_help_tab(self):
        tab = self.help_tab
        tab.columnconfigure(0, weight=1)
        ttk.Label(tab, text=f"{APP_NAME} {APP_VERSION}", style="HelpTitle.TLabel").grid(row=0, column=0, sticky="w", pady=(0, 12))
        help_text = (
            "Factura individual: selecciona una factura, el informe de Alma Analytics, la plantilla y el destino.\n\n"
            "Modo lote: todas las facturas deben pertenecer al mismo proveedor y biblioteca. La lista permite anadir, quitar, ordenar y detectar duplicados.\n\n"
            "Informe de validacion: se genera cuando faltan PO Lines, hay coincidencias ambiguas, duplicados o diferencias entre Quantity y Quantity for Pricing. "
            "Una diferencia de cantidad requiere revision manual y no determina por si sola si la facturacion es parcial o total.\n\n"
            "Limitaciones confirmadas por Ex Libris: la plantilla Excel no permite indicar Line Exclusive ni declarar explicitamente facturacion parcial o completa.\n\n"
            "Licencia: GPL-3.0-only."
        )
        ttk.Label(tab, text=help_text, wraplength=940, justify="left", style="App.TLabel").grid(row=1, column=0, sticky="nw")
        links = ttk.LabelFrame(tab, text="Documentacion y soporte", padding=14, style="App.TLabelframe")
        links.grid(row=2, column=0, sticky="ew", pady=(20, 0))
        ttk.Button(links, text="Abrir repositorio de GitHub", command=lambda: self._open_url(PROJECT_REPOSITORY_URL)).pack(anchor="w", pady=4)
        ttk.Button(links, text="Abrir Biblioguia / LibGuides", command=lambda: self._open_url(PROJECT_GUIDE_URL)).pack(anchor="w", pady=4)
        ttk.Label(links, text=PROJECT_REPOSITORY_URL, style="App.TLabel").pack(anchor="w", pady=(10, 2))
        ttk.Label(links, text=PROJECT_GUIDE_URL, style="App.TLabel").pack(anchor="w")

    def select_logo(self):
        current = self.logo_path_var.get().strip()
        initial = os.path.dirname(current) if current and os.path.isfile(current) else self._application_dir()
        path = filedialog.askopenfilename(
            title="Seleccionar logo institucional",
            initialdir=initial,
            filetypes=[
                ("Imagenes compatibles", "*.png *.jpg *.jpeg"),
                ("Imagen PNG", "*.png"),
                ("Imagen JPEG", "*.jpg *.jpeg"),
                ("Todos los archivos", "*.*"),
            ],
        )
        if not path:
            return
        self.logo_path_var.set(path)
        if self._apply_logo(path, show_warning=True):
            self.logo_status_label.config(text=f"Logo cargado: {os.path.basename(path)}")

    def remove_logo(self):
        self.logo_path_var.set("")
        self.logo_photo = None
        self.logo_label.config(image="")
        self.logo_status_label.config(text="Sin logo. PNG, JPG o JPEG; altura maxima: 100 px.")

    def _load_logo_from_preferences(self, show_warning=False):
        path = self.logo_path_var.get().strip()
        if not path:
            self.logo_photo = None
            if hasattr(self, "logo_label"):
                self.logo_label.config(image="")
            return False
        return self._apply_logo(path, show_warning=show_warning)

    def _apply_logo(self, path, show_warning=False):
        if not hasattr(self, "logo_label"):
            return False
        if not os.path.isfile(path):
            self.logo_photo = None
            self.logo_label.config(image="")
            if hasattr(self, "logo_status_label"):
                self.logo_status_label.config(text="El logo configurado no esta disponible.")
            if show_warning:
                messagebox.showwarning("Logo no disponible", f"No se encontro el archivo de logo:\n{path}")
            return False
        try:
            with Image.open(path) as opened:
                image = opened.copy()
            if image.mode not in ("RGB", "RGBA", "L", "LA"):
                image = image.convert("RGBA" if "transparency" in image.info else "RGB")
            if image.height > MAX_LOGO_HEIGHT:
                ratio = MAX_LOGO_HEIGHT / image.height
                new_width = max(1, round(image.width * ratio))
                image = image.resize((new_width, MAX_LOGO_HEIGHT), Image.Resampling.LANCZOS)
            self.logo_photo = ImageTk.PhotoImage(image)
            self.logo_label.config(image=self.logo_photo)
            return True
        except (UnidentifiedImageError, OSError, ValueError) as e:
            logging.warning(f"No se pudo cargar el logo {path}: {e}")
            self.logo_photo = None
            self.logo_label.config(image="")
            if hasattr(self, "logo_status_label"):
                self.logo_status_label.config(text="El archivo seleccionado no es una imagen valida.")
            if show_warning:
                messagebox.showwarning("Logo no valido", f"No se pudo cargar el logo:\n{path}\n\n{e}")
            return False

    def _choose_preference_dir(self, variable):
        initial = self._valid_dir(variable.get()) or self._application_dir()
        folder = filedialog.askdirectory(initialdir=initial, title="Seleccionar carpeta")
        if folder:
            variable.set(folder)

    def save_preferences(self):
        institution = self.institution_var.get().strip() or self.DEFAULT_INSTITUTION
        self.institution_var.set(institution)
        for value, label in [
            (self.custom_invoice_dir.get().strip(), "Directorio de facturas FACe"),
            (self.custom_alma_dir.get().strip(), "Directorio de Alma"),
            (self.custom_template_dir.get().strip(), "Directorio de plantilla"),
            (self.custom_output_dir.get().strip(), "Directorio de salida"),
        ]:
            if value and not os.path.isdir(value):
                messagebox.showerror("Preferencias no validas", f"{label} no es una carpeta valida:\n{value}")
                return
        logo_path = self.logo_path_var.get().strip()
        if logo_path and not self._apply_logo(logo_path, show_warning=True):
            return
        self.preferences.update({
            "institution": institution,
            "custom_invoice_dir": self.custom_invoice_dir.get().strip(),
            "custom_alma_dir": self.custom_alma_dir.get().strip(),
            "custom_template_dir": self.custom_template_dir.get().strip(),
            "custom_output_dir": self.custom_output_dir.get().strip(),
            "logo_path": logo_path,
        })
        try:
            self._save_preferences_file()
            self._update_title()
            self._load_logo_from_preferences(show_warning=False)
            messagebox.showinfo("Preferencias", "Las preferencias se han guardado correctamente.")
        except ValueError as e:
            messagebox.showerror("Error al guardar", str(e))

    def reset_preferences(self):
        if not messagebox.askyesno("Restablecer preferencias", "¿Quieres restablecer la institucion y todos los directorios personalizados?"):
            return
        defaults = self._default_preferences()
        self.preferences = defaults
        self.institution_var.set(defaults["institution"])
        self.custom_invoice_dir.set("")
        self.custom_alma_dir.set("")
        self.custom_template_dir.set("")
        self.custom_output_dir.set("")
        self.logo_path_var.set("")
        self.logo_photo = None
        self.logo_label.config(image="")
        self.logo_status_label.config(text="Sin logo. PNG, JPG o JPEG; altura maxima: 100 px.")
        self.template_file.set("")
        try:
            self._save_preferences_file()
        except ValueError as e:
            messagebox.showerror("Error al restablecer", str(e))
            return
        self._update_title()
        messagebox.showinfo("Preferencias", "Se han restablecido los valores iniciales.")

    def _open_url(self, url):
        try:
            webbrowser.open_new(url)
        except Exception as e:
            messagebox.showerror("Error al abrir enlace", f"No se pudo abrir el enlace:\n{url}\n\n{e}")

    def _open_folder(self, folder_path):
        folder_path = os.path.realpath(os.path.abspath(folder_path))
        if not os.path.isdir(folder_path):
            messagebox.showwarning("Aviso", "No se encontro la carpeta de salida.")
            return
        try:
            system_name = platform.system()
            if system_name == "Windows":
                os.startfile(folder_path)  # nosec B606 - ruta normalizada y validada como directorio
                return

            command_name = "open" if system_name == "Darwin" else "xdg-open"
            opener_path = shutil.which(command_name)
            if not opener_path or not os.path.isabs(opener_path):
                raise FileNotFoundError(
                    f"No se encontro una ruta absoluta para la utilidad '{command_name}'."
                )

            subprocess.run(  # nosec B603 - ejecutable absoluto y argumento validado
                [opener_path, folder_path],
                check=False,
                shell=False,
            )
        except (OSError, subprocess.SubprocessError) as e:
            messagebox.showerror("Error al abrir carpeta", f"No se pudo abrir la carpeta:\n{folder_path}\n\n{e}")

    def open_output_folder(self):
        output_path = self.output_file.get().strip()
        folder = os.path.dirname(output_path) if output_path else ""
        if not folder:
            messagebox.showwarning("Aviso", "No hay una carpeta de salida definida.")
            return
        self._open_folder(folder)

    def on_mode_change(self, initial=False):
        mode = self.conversion_mode.get()
        if not initial:
            self.invoice_files = []
            self.invoice_metadata = {}
            self.invoice_file.set("")
            self.output_file.set("")
            self._refresh_batch_tree()
        if mode == "batch":
            self.invoice_label.config(text="Facturas XSIG / XML / TXT del mismo proveedor y biblioteca *")
            self.invoice_button.config(text="Seleccionar facturas...")
            self.batch_frame.grid()
            self.status_text.set("Estado: modo lote. Selecciona facturas del mismo proveedor y biblioteca.")
        else:
            self.invoice_label.config(text="Factura electronica XSIG / XML / TXT *")
            self.invoice_button.config(text="Seleccionar factura...")
            self.batch_frame.grid_remove()
            self.status_text.set("Estado: modo factura individual.")

    def select_invoice_or_batch(self):
        if self.conversion_mode.get() == "batch":
            self.add_batch_files(replace=True)
        else:
            self.select_invoice_single()

    def _invoice_filetypes(self):
        return [("Facturas Universitas XXI / Facturae", "*.xsig *.xml *.txt"), ("Archivos XSIG", "*.xsig"), ("Archivos XML", "*.xml"), ("Archivos TXT", "*.txt"), ("Todos los archivos", "*.*")]

    def select_invoice_single(self):
        initial = self._initial_dir("custom_invoice_dir", "last_invoice_dir")
        path = filedialog.askopenfilename(title="Seleccionar factura XSIG / XML / TXT", initialdir=initial, filetypes=self._invoice_filetypes())
        if not path:
            return
        try:
            metadata = self._read_invoice_metadata(path)
        except Exception as e:
            messagebox.showerror("Factura no valida", str(e))
            return
        self.invoice_files = [path]
        self.invoice_metadata = {path: metadata}
        self.invoice_file.set(path)
        self._remember_dir("last_invoice_dir", path)
        self._set_default_output(path, batch=False)
        self.status_text.set(f"Estado: factura {metadata['number']} seleccionada. Proveedor: {metadata['vendor']}.")

    def add_batch_files(self, replace=False):
        initial = self._initial_dir("custom_invoice_dir", "last_invoice_dir")
        paths = filedialog.askopenfilenames(title="Seleccionar facturas del mismo proveedor y biblioteca", initialdir=initial, filetypes=self._invoice_filetypes())
        if not paths:
            return
        if replace:
            self.invoice_files = []
            self.invoice_metadata = {}
        existing_norm = {os.path.normcase(os.path.abspath(p)) for p in self.invoice_files}
        existing_numbers = {m["number"] for m in self.invoice_metadata.values() if m.get("number")}
        added = 0
        warnings = []
        for path in paths:
            norm = os.path.normcase(os.path.abspath(path))
            if norm in existing_norm:
                warnings.append(f"Ruta repetida: {os.path.basename(path)}")
                continue
            try:
                metadata = self._read_invoice_metadata(path)
            except Exception as e:
                warnings.append(f"No se pudo leer {os.path.basename(path)}: {e}")
                continue
            number = metadata.get("number")
            if number and number in existing_numbers:
                warnings.append(f"Numero de factura duplicado {number}: {os.path.basename(path)}")
                continue
            self.invoice_files.append(path)
            self.invoice_metadata[path] = metadata
            existing_norm.add(norm)
            if number:
                existing_numbers.add(number)
            added += 1
        self._refresh_batch_tree()
        if self.invoice_files:
            self.invoice_file.set(f"{len(self.invoice_files)} facturas seleccionadas")
            self._remember_dir("last_invoice_dir", self.invoice_files[0])
            self._set_default_output(self.invoice_files[0], batch=True)
        self.status_text.set(f"Estado: {len(self.invoice_files)} facturas en el lote.")
        if warnings:
            messagebox.showwarning("Facturas no incorporadas", "\n".join(warnings))
        elif added == 0:
            messagebox.showinfo("Lote", "No se incorporaron facturas nuevas.")

    def select_invoice_batch(self):
        self.add_batch_files(replace=True)

    def _read_invoice_metadata(self, path):
        parsed = FacturaeParser(path).parse()
        if len(parsed) != 1:
            raise ValueError(f"El fichero debe contener una sola factura y contiene {len(parsed)}.")
        inv = parsed[0]
        return {"number": inv.invoice_number or "(sin numero)", "vendor": inv.vendor_account or inv.seller_name or "Proveedor no identificado"}

    def _refresh_batch_tree(self):
        if not hasattr(self, "batch_tree"):
            return
        for item in self.batch_tree.get_children():
            self.batch_tree.delete(item)
        for path in self.invoice_files:
            meta = self.invoice_metadata.get(path, {})
            self.batch_tree.insert("", "end", iid=path, values=(meta.get("number", ""), meta.get("vendor", ""), os.path.basename(path)))

    def remove_selected_batch_files(self):
        selected = list(self.batch_tree.selection())
        if not selected:
            messagebox.showinfo("Lote", "Selecciona al menos una factura de la lista.")
            return
        selected_set = set(selected)
        self.invoice_files = [p for p in self.invoice_files if p not in selected_set]
        for p in selected:
            self.invoice_metadata.pop(p, None)
        self._refresh_batch_tree()
        self.invoice_file.set(f"{len(self.invoice_files)} facturas seleccionadas" if self.invoice_files else "")
        self.status_text.set(f"Estado: {len(self.invoice_files)} facturas en el lote.")

    def clear_batch_files(self):
        if self.invoice_files and not messagebox.askyesno("Vaciar lote", "¿Quieres eliminar todas las facturas del lote?"):
            return
        self.invoice_files = []
        self.invoice_metadata = {}
        self.invoice_file.set("")
        self.output_file.set("")
        self._refresh_batch_tree()
        self.status_text.set("Estado: lote vacio.")

    def sort_batch_files(self):
        self.invoice_files.sort(key=lambda p: (self.invoice_metadata.get(p, {}).get("number", "").casefold(), os.path.basename(p).casefold()))
        self._refresh_batch_tree()
        self.status_text.set("Estado: lote ordenado por numero de factura.")

    def _set_default_output(self, invoice_path, batch):
        custom_output = self._valid_dir(self.preferences.get("custom_output_dir", ""))
        invoice_folder = os.path.dirname(invoice_path)
        output_folder = custom_output or invoice_folder or self._initial_dir("custom_output_dir", "last_output_dir")
        base = os.path.basename(os.path.splitext(invoice_path)[0])
        if batch:
            output_name = f"Lote_{base}_Alma.xlsx"
        else:
            output_name = f"{base}_Alma.xlsx"
        self.output_file.set(os.path.join(output_folder, output_name))

    def select_alma_pol(self):
        fallback = os.path.dirname(self.invoice_files[0]) if self.invoice_files else None
        initial = self._initial_dir("custom_alma_dir", "last_alma_dir", fallback)
        path = filedialog.askopenfilename(title="Seleccionar fichero Excel de Alma Analytics", initialdir=initial, filetypes=[("Archivos Excel", "*.xlsx"), ("Todos los archivos", "*.*")])
        if path:
            self.alma_pol_file.set(path)
            self._remember_dir("last_alma_dir", path)
            self.status_text.set("Estado: fichero Excel de Alma Analytics seleccionado.")

    def select_template(self):
        fallback = os.path.dirname(self.alma_pol_file.get()) if self.alma_pol_file.get() else None
        initial = self._initial_dir("custom_template_dir", "last_template_dir", fallback)
        path = filedialog.askopenfilename(title="Seleccionar plantilla Alma .xlsx", initialdir=initial, filetypes=[("Archivos Excel", "*.xlsx"), ("Todos los archivos", "*.*")])
        if path:
            self.template_file.set(path)
            self.preferences["last_template_path"] = path
            self._remember_dir("last_template_dir", path)
            self.status_text.set("Estado: plantilla Alma seleccionada.")

    def select_output(self):
        fallback = os.path.dirname(self.invoice_files[0]) if self.invoice_files else None
        initial = self._initial_dir("custom_output_dir", "last_output_dir", fallback)
        initialfile = os.path.basename(self.output_file.get()) if self.output_file.get() else "FacturaeToAlma.xlsx"
        path = filedialog.asksaveasfilename(title="Definir archivo Excel de salida", initialdir=initial, initialfile=initialfile, defaultextension=".xlsx", filetypes=[("Archivos Excel", "*.xlsx"), ("Todos los archivos", "*.*")])
        if path:
            self.output_file.set(path)
            self._remember_dir("last_output_dir", path)
            self.status_text.set("Estado: archivo de salida definido.")

    def clear_fields(self):
        self.invoice_files = []
        self.invoice_metadata = {}
        self.invoice_file.set("")
        self.output_file.set("")
        self._refresh_batch_tree()
        if self.alma_pol_file.get().strip() and messagebox.askyesno("Limpiar fichero Excel de Alma", "¿Quieres limpiar tambien el fichero Excel de Alma Analytics?\n\nLa plantilla Alma se conservara."):
            self.alma_pol_file.set("")
        self.status_text.set("Estado: datos limpiados. La plantilla Alma se conserva.")

    def _validate_inputs(self):
        if not self.invoice_files:
            raise ValueError("Debes seleccionar una factura o un lote de facturas.")
        allowed = {".xsig", ".xml", ".txt"}
        for path in self.invoice_files:
            if os.path.splitext(path)[1].lower() not in allowed:
                raise ValueError(f"Extension de factura no valida: {path}")
            if not os.path.isfile(path):
                raise FileNotFoundError(f"No existe la factura: {path}")
        alma = self.alma_pol_file.get().strip()
        template = self.template_file.get().strip()
        output = self.output_file.get().strip()
        if not alma:
            raise ValueError("Debes seleccionar el fichero Excel de Alma Analytics.")
        if not template:
            raise ValueError("Debes seleccionar la plantilla Alma.")
        if not output:
            raise ValueError("Debes definir el Excel de salida.")
        if not os.path.isfile(alma):
            raise FileNotFoundError(f"No existe el fichero Excel de Alma Analytics: {alma}")
        if not os.path.isfile(template):
            raise FileNotFoundError(f"No existe la plantilla Alma: {template}")
        duplicate_numbers = {}
        for path in self.invoice_files:
            number = self.invoice_metadata.get(path, {}).get("number")
            if number:
                duplicate_numbers.setdefault(number, []).append(path)
        repeated = {n: paths for n, paths in duplicate_numbers.items() if len(paths) > 1}
        if repeated:
            raise ValueError("Hay numeros de factura duplicados en el lote: " + ", ".join(sorted(repeated)))
        return list(self.invoice_files), alma, output, template

    def run_conversion(self):
        if self.worker_thread and self.worker_thread.is_alive():
            messagebox.showinfo("Conversion en curso", "Ya hay una conversion en curso.")
            return
        try:
            inputs = self._validate_inputs()
        except (ValueError, FileNotFoundError) as e:
            messagebox.showerror("Datos de entrada no validos", str(e))
            self.status_text.set("Estado: revisa los datos de entrada.")
            return
        self.convert_button.config(state="disabled")
        self.status_text.set("Estado: convirtiendo en segundo plano...")
        self.worker_thread = threading.Thread(target=self._conversion_worker, args=inputs, daemon=True)
        self.worker_thread.start()
        self.root.after(100, self._poll_worker)

    def _conversion_worker(self, invoice_paths, alma_path, output_path, template_path):
        try:
            result = convert_facturae_to_alma_excel(invoice_paths, alma_path, output_path, template_path)
            self.worker_queue.put(("success", result, output_path))
        except Exception as e:
            logging.exception("Error durante la conversion")
            self.worker_queue.put(("error", e, traceback.format_exc()))

    def _poll_worker(self):
        try:
            kind, payload, extra = self.worker_queue.get_nowait()
        except queue.Empty:
            if self.worker_thread and self.worker_thread.is_alive():
                self.root.after(100, self._poll_worker)
            return
        self.convert_button.config(state="normal")
        if kind == "success":
            self._handle_conversion_success(payload, extra)
        else:
            self._handle_conversion_error(payload, extra)

    def _build_summary_message(self, summary):
        return (
            "\n\nResumen de lineas:\n"
            f"- Lineas procesadas: {summary['total_lines']}\n"
            f"- Coincidencias por ISBN: {summary['match_isbn']}\n"
            f"- Coincidencias por ISSN: {summary['match_issn']}\n"
            f"- Coincidencias por titulo: {summary['match_title']}\n"
            f"- Duplicados ISBN: {summary['duplicate_isbn']}\n"
            f"- Duplicados ISSN: {summary['duplicate_issn']}\n"
            f"- Titulos ambiguos: {summary['ambiguous_title']}\n"
            f"- Sin coincidencia: {summary['no_match']}\n"
            f"- Lineas sin PO Line: {summary['lines_without_po_line']}"
        )

    def _handle_conversion_success(self, result, output_path):
        total_invoices, total_issues, validation_path, summary = result
        heading = "CONVERSION COMPLETA SIN INCIDENCIAS" if not validation_path else "INCIDENCIAS EN LA CONVERSION"
        message = (
            f"{heading}\n\n"
            f"Facturas procesadas: {total_invoices}\n"
            f"Avisos de validacion: {total_issues}\n\n"
            f"Excel Alma generado:\n{output_path}"
        ) + self._build_summary_message(summary)
        if validation_path:
            message += f"\n\nInforme de validacion:\n{validation_path}"
        if summary.get("lines_without_po_line", 0) > 0:
            message += "\n\nAtencion: hay lineas sin PO Line. Revisa el informe antes de cargar en Alma."
        self._remember_dir("last_output_dir", output_path)
        self.status_text.set("Estado: conversion completada con incidencias." if validation_path else "Estado: conversion completada sin incidencias.")
        if validation_path:
            messagebox.showwarning("Incidencias en la conversion", message)
        else:
            messagebox.showinfo("Conversion completada", message)

    def _handle_conversion_error(self, error, trace):
        if isinstance(error, DefusedXmlException):
            title = "XML bloqueado por seguridad"
        elif isinstance(error, PermissionError):
            title = "Fichero bloqueado o sin permisos"
        elif isinstance(error, FileNotFoundError):
            title = "Fichero no encontrado"
        elif isinstance(error, ValueError):
            title = "Datos o formato no validos"
        else:
            title = "Error durante la conversion"
        logging.error(trace)
        self.status_text.set("Estado: error durante la conversion.")
        messagebox.showerror(title, f"{error}\n\nConsulta el archivo de log '{LOG_FILE}' para mas detalles.")

    def _on_close(self):
        if self.worker_thread and self.worker_thread.is_alive():
            if not messagebox.askyesno("Conversion en curso", "Hay una conversion en curso. ¿Quieres cerrar la aplicacion de todos modos?"):
                return
        self.root.destroy()

if __name__ == "__main__":
    root = tk.Tk()
    app = FacturaeApp(root)
    root.mainloop()
